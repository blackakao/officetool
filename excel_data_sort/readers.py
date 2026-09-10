from pathlib import Path

from .models import SheetData, WorkbookData


# Guard against worksheets whose formatting extends to Excel's entire grid.
MAX_SHEET_CELLS = 2_000_000


def _check_size(name, rows, columns):
    if rows * columns > MAX_SHEET_CELLS:
        raise ValueError(f"{name}: 읽기 범위가 {MAX_SHEET_CELLS:,}셀을 초과합니다. 불필요한 끝 행·열을 삭제해 주세요.")


def read_workbook(path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        return _read_xlsx(path)
    if path.suffix.lower() == ".xls":
        return _read_xls(path)
    raise ValueError(".xlsx 또는 .xls 파일만 지원합니다.")


def _read_xlsx(path):
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=True)
    try:
        sheets = []
        for sheet in book.worksheets:
            _check_size(sheet.title, sheet.max_row, sheet.max_column)
            hidden_columns = set()
            for dimension in sheet.column_dimensions.values():
                if dimension.hidden:
                    hidden_columns.update(range(dimension.min, dimension.max + 1))
            sheets.append(SheetData(
                name=sheet.title,
                rows=[list(row) for row in sheet.iter_rows(values_only=True)],
                merges=[(area.min_row, area.max_row, area.min_col, area.max_col)
                        for area in sheet.merged_cells.ranges],
                hidden_rows={index for index, dimension in sheet.row_dimensions.items() if dimension.hidden},
                hidden_columns=hidden_columns,
                hidden=sheet.sheet_state != "visible",
            ))
        return WorkbookData(path, sheets, [
            "수식은 Excel에 마지막으로 저장된 계산값을 읽습니다. 값이 비어 있으면 Excel에서 재계산 후 저장해 주세요."
        ])
    finally:
        book.close()


def _read_xls(path):
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError("XLS 읽기에는 xlrd가 필요합니다. requirements.txt의 패키지를 설치해 주세요.") from exc

    book = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
    try:
        sheets = []
        for sheet in book.sheets():
            _check_size(sheet.name, sheet.nrows, sheet.ncols)
            rows = []
            for row_index in range(sheet.nrows):
                values = []
                for cell in sheet.row(row_index):
                    value = cell.value
                    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        value = None
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        converted = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                        value = converted.time() if 0 <= value < 1 else converted
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "#ERROR!")
                    values.append(value)
                rows.append(values)
            sheets.append(SheetData(
                name=sheet.name, rows=rows,
                merges=[(r0 + 1, r1, c0 + 1, c1) for r0, r1, c0, c1 in sheet.merged_cells],
                hidden_rows={index + 1 for index in range(sheet.nrows)
                             if getattr(sheet.rowinfo_map.get(index), "hidden", sheet.default_row_hidden)},
                hidden_columns={index + 1 for index, info in sheet.colinfo_map.items() if info.hidden},
                hidden=bool(sheet.visibility),
            ))
        return WorkbookData(path, sheets)
    finally:
        book.release_resources()
