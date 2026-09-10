from pathlib import Path
import re
import tempfile

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


def export_excel(datasets, output_path, source_paths=()):
    """Write values (never executable formulas) to an atomic XLSX output."""
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("출력 파일 확장자는 .xlsx여야 합니다.")
    if output_path.resolve() in {Path(path).resolve() for path in source_paths}:
        raise ValueError("원본 Excel 파일에는 덮어쓸 수 없습니다. 다른 저장 경로를 선택해 주세요.")
    if not datasets or not any(dataset.rows for dataset in datasets):
        raise ValueError("저장할 Dataset이 없습니다.")
    book = Workbook()
    book.remove(book.active)
    titles = set()
    temporary = None
    try:
        for dataset in datasets:
            if len(dataset.rows) > 1_048_575 or len(dataset.columns) > 16_384:
                raise ValueError("Excel의 행·열 한도를 초과했습니다. 시트 또는 테이블을 분리해 주세요.")
            base = re.sub(r"[\\/*?:\[\]]", "_", dataset.name).strip("'")[:31] or "Dataset"
            title, number = base, 1
            while title.casefold() in titles:
                number += 1
                suffix = f"_{number}"
                title = base[:31 - len(suffix)] + suffix
            titles.add(title.casefold())
            sheet = book.create_sheet(title)
            for row_index, values in enumerate([dataset.columns, *dataset.rows], 1):
                for column_index, value in enumerate(values, 1):
                    if isinstance(value, str) and (len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value)):
                        raise ValueError(f"{title} {row_index}행: Excel에 저장할 수 없는 문자열이 있습니다.")
                    cell = sheet.cell(row_index, column_index, value)
                    if isinstance(value, str):
                        cell.data_type = "s"
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        # Temporary file resides beside the destination so replace is atomic.
        with tempfile.NamedTemporaryFile(dir=output_path.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        temporary.replace(output_path)
        temporary = None
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
