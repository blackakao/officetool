import copy
import json
import os

import pytest

from business_planning.minutes import (
    MinutesStore, atomic_json, export_docx, new_agenda, new_meeting,
    parse_response, validate,
)


@pytest.fixture
def meeting():
    m = new_meeting()
    m.update(version=1, branch_id='branch-1', branch_name='테스트 지점', title='운영회의', place='회의실')
    m['attendees'] = [dict(id='p1', name='위원1', position='시설장', role='진행', present=True),
                      dict(id='p2', name='위원2', position='대표', role='위원', present=True),
                      dict(id='p3', name='위원3', position='대표', role='위원', present=False)]
    for title in ['사업계획', '예산편성']:
        a = new_agenda()
        a.update(title=title, notes='복리후생비 100원 검토', final='검토 중입니다.')
        m['agendas'].append(a)
    return m


def response(meeting):
    return json.dumps({'agendas': [dict(id=a['id'], speeches=[
        dict(attendee_id='p1', text='복리후생비 100원을 검토하겠습니다.'),
        dict(attendee_id='p2', text='검토할 자료가 있습니까?')]) for a in meeting['agendas']]})


def test_roundtrip_and_branch_history(tmp_path, meeting):
    store = MinutesStore(tmp_path)
    store.save(meeting)
    restored = store.load(meeting['id'])
    assert restored['agendas'] == meeting['agendas']
    assert restored['attendees'] == meeting['attendees']
    assert store.history('other') == []
    assert store.history('branch-1')[0]['id'] == meeting['id']
    store.roster('branch-1', meeting['attendees'])
    assert store.roster('branch-1') == meeting['attendees']
    assert store.roster('other') == []
    with pytest.raises(ValueError):
        store.load('../outside')


def test_atomic_failure_preserves_original(tmp_path, monkeypatch):
    path = tmp_path / 'data.json'
    atomic_json(path, {'original': True})
    original = path.read_bytes()
    def fail(*_):
        raise OSError('disk error')
    monkeypatch.setattr(os, 'replace', fail)
    with pytest.raises(OSError):
        atomic_json(path, {'original': False})
    assert path.read_bytes() == original
    assert not list(tmp_path.glob('*.tmp'))


def test_corrupt_storage_is_not_reset(tmp_path):
    store = MinutesStore(tmp_path)
    path = store.directory / 'rosters.json'
    path.parent.mkdir(parents=True)
    path.write_text('{broken', encoding='utf-8')
    with pytest.raises(json.JSONDecodeError):
        store.roster('b', [])
    assert path.read_text(encoding='utf-8') == '{broken'


def test_response_rejects_missing_agenda_and_absent_speaker(meeting):
    assert len(parse_response(response(meeting), meeting)) == 2
    data = json.loads(response(meeting))
    data['agendas'][0]['speeches'][0]['attendee_id'] = 'p3'
    with pytest.raises(ValueError):
        parse_response(json.dumps(data), meeting)
    with pytest.raises(ValueError):
        parse_response('{"agendas": []}', meeting)
    data = json.loads(response(meeting))
    data['agendas'][1]['id'] = data['agendas'][0]['id']
    with pytest.raises(ValueError):
        parse_response(json.dumps(data), meeting)


def test_validation_and_docx_blank_signatures(tmp_path, meeting):
    from docx import Document
    validate(meeting)
    path = tmp_path / 'minutes.docx'
    with pytest.raises(ValueError, match='확인'):
        export_docx(meeting, path)
    meeting['confirmed'] = True
    export_docx(meeting, path)
    doc = Document(path)
    assert doc.tables[0].rows[0].cells[1].text == '운영회의'
    signatures = doc.tables[-1]
    assert len(signatures.rows) == 3
    assert all(row.cells[3].text == '' for row in signatures.rows[1:])
    assert not doc.inline_shapes
    original = path.read_bytes()
    with pytest.raises(FileExistsError):
        export_docx(meeting, path)
    assert path.read_bytes() == original
    meeting['end'] = '09:00'
    with pytest.raises(ValueError, match='종료'):
        validate(meeting)


@pytest.fixture
def app():
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def test_page_edits_roundtrip_and_confirmation_reset(app, tmp_path, meeting, monkeypatch):
    from ui.pages import committee_minutes_page as page_module
    path = tmp_path / 'branches.json'
    path.write_text(json.dumps([dict(branch_name='테스트 지점', organization_code='branch-1')]), encoding='utf-8')
    monkeypatch.setattr(page_module, 'user_data_path', lambda _: path)
    page = page_module.CommitteeMinutesPage(MinutesStore(tmp_path))
    page.display(meeting)
    assert page.prompt.toPlainText()
    assert '검토 중입니다.' in page.final_text.toPlainText()
    page.prompt.setPlainText('위원1은 진행자, 위원2는 질문을 합니다.')
    page.final_text.setPlainText('사용자 수정')
    page.quarter.setCurrentIndex(2)
    page.confirmed.setChecked(True)
    page.save()
    restored = page.store.load(meeting['id'])
    assert restored['final'] == '사용자 수정'
    assert restored['quarter'] == 3
    assert restored['legacy_source']['agendas'] == meeting['agendas']
    assert restored['prompt'] == '위원1은 진행자, 위원2는 질문을 합니다.'
    assert restored['confirmed']
    page.title.setText('주제 변경')
    assert not page.confirmed.isChecked()
    page.display(restored)
    assert page.collect()['attendees'] == meeting['attendees']
    page.apply_result('위원1: 안녕하세요.\n위원2: 감사합니다.')
    assert page.collect()['draft'] == page.collect()['final']
    assert not page.confirmed.isChecked()
    page.deleteLater()


def test_menu_pages_and_existing_navigation(app, monkeypatch):
    import main
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QWidget
    from ui.pages.committee_minutes_page import PlannedPage
    names = ['BranchPage', 'LoginTool', 'LoginLogPage', 'UnifiedDocumentTool',
             'TableListPage', 'FederationTool', 'CareforTool', 'LaborCostRatioPage',
             'ServiceFeePage', 'AnnualLeavePage', 'MonthlyWorkCheckPage', 'MemoPage',
             'IncompleteTaskPage', 'BranchTaskSettingsPage', 'AttendanceAnalysisPage',
             'ExcelDataSortPage', 'CommitteeMinutesPage']
    for name in names:
        monkeypatch.setattr(main, name, QWidget)
    window = main.MainWindow()
    entries = {window.menu.item(i).text().strip(): window.menu.item(i)
               for i in range(window.menu.count())}
    assert entries['사업기획'].data(Qt.UserRole) is None
    assert isinstance(entries['예산편성 (예정)'].data(Qt.UserRole), PlannedPage)
    for label in ['사업계획 (예정)', '운영위원회 회의록 작성', '로그인툴', '문서작성툴']:
        window.on_menu_clicked(entries[label])
        assert window.stacked_widget.currentWidget() is entries[label].data(Qt.UserRole)
    window.deleteLater()


def test_worker_failure_keeps_meeting(app, meeting, monkeypatch):
    from ui.pages import committee_minutes_page as module
    original = copy.deepcopy(meeting)
    def fail(*args):
        # A provider must never be allowed to mutate the UI's source data.
        args[0]['title'] = 'changed by provider'
        raise ValueError('AI 사용량 제한')
    monkeypatch.setattr(module, 'generate_local', fail)
    worker = module.MinutesWorker(meeting, {'directory': 'unused'})
    failures = []
    worker.failed.connect(failures.append)
    worker.run()
    assert failures == ['AI 사용량 제한']
    assert meeting == original


def test_local_generation_isolates_agendas_and_keeps_inputs(meeting, tmp_path):
    from business_planning.local_ai import generate_local
    original = copy.deepcopy(meeting)
    calls = []
    def completion(tokenizer, model, messages, cancelled):
        payload = json.loads(messages[1]['content'])
        calls.append(payload)
        assert len(payload['agendas']) == 1
        assert 'final' not in payload['agendas'][0]
        assert len(payload['attendees']) == 2
        return json.dumps({'agendas': [{'id': 'a1', 'speeches': [
            {'attendee_id': 'p1', 'text': '복리후생비 100원을 검토하겠습니다.'},
            {'attendee_id': 'p2', 'text': '검토 자료를 확인하겠습니다.'}]}]})
    result = generate_local(meeting, tmp_path, loader=lambda _: (None, None), completion=completion)
    assert set(result) == {a['id'] for a in meeting['agendas']}
    assert len(calls) == 2
    assert meeting == original


def test_local_failure_and_cancel_do_not_apply_partial_result(meeting, tmp_path):
    from business_planning.local_ai import generate_local, load_engine
    original = copy.deepcopy(meeting)
    with pytest.raises(ValueError, match='다운로드'):
        load_engine(tmp_path)
    with pytest.raises(ValueError, match='응답 형식'):
        generate_local(meeting, tmp_path, loader=lambda _: (None, None), completion=lambda *_: 'not json')
    with pytest.raises(ValueError, match='취소'):
        generate_local(meeting, tmp_path, loader=lambda _: (None, None), cancelled=lambda: True)
    assert meeting == original


def test_external_settings_are_ignored(app, tmp_path, monkeypatch):
    from ui.pages import committee_minutes_page as module
    branches = tmp_path / 'branches.json'
    branches.write_text('[]', encoding='utf-8')
    monkeypatch.setattr(module, 'user_data_path', lambda _: branches)
    store = MinutesStore(tmp_path)
    atomic_json(store.directory / 'ai_settings.json', {'endpoint': 'https://external.invalid', 'model': 'old'})
    page = module.CommitteeMinutesPage(store)
    assert set(page.settings()) == {'directory'}
    assert page.settings()['directory'].is_relative_to(store.directory)
    page.deleteLater()


def test_model_loader_never_uses_remote_code_or_network(tmp_path, monkeypatch):
    import sys
    from types import SimpleNamespace
    from business_planning.local_ai import MODEL_FILES, load_engine
    for name in MODEL_FILES:
        (tmp_path / name).write_text('fixture', encoding='utf-8')
    calls = []
    def load(path, **kwargs):
        calls.append(kwargs)
        return SimpleNamespace(eval=lambda: None)
    monkeypatch.setitem(sys.modules, 'transformers', SimpleNamespace(
        AutoTokenizer=SimpleNamespace(from_pretrained=load),
        AutoModelForCausalLM=SimpleNamespace(from_pretrained=load)))
    load_engine(tmp_path)
    assert all(c['local_files_only'] and not c['trust_remote_code'] for c in calls)
    assert calls[1]['use_safetensors']


def test_prompt_ui_only_has_simplified_inputs(app, tmp_path, monkeypatch):
    from ui.pages import committee_minutes_page as module
    branches = tmp_path / 'branches.json'
    branches.write_text('[]', encoding='utf-8')
    monkeypatch.setattr(module, 'user_data_path', lambda _: branches)
    page = module.CommitteeMinutesPage(MinutesStore(tmp_path))
    assert [page.tabs.tabText(i) for i in range(page.tabs.count())] == ['참석자', '프롬프트 / 회의록']
    assert not hasattr(page, 'year') and not hasattr(page, 'session')
    assert not hasattr(page, 'agenda_fields') and not hasattr(page, 'opening')
    assert [page.quarter.itemText(i) for i in range(4)] == ['1분기', '2분기', '3분기', '4분기']
    page.deleteLater()


def test_prompt_generation_obeys_roles_and_excludes_absentees(meeting, tmp_path):
    from business_planning.local_ai import generate_local
    from business_planning.minutes import prompt_meeting
    m = prompt_meeting(meeting)
    m['prompt'] = '위원1은 공손한 진행자, 위원2는 질문과 감사를 한다. 2026년 예산 적법성 검토.'
    original = copy.deepcopy(m)
    def complete(*args):
        prompt = args[2][1]['content']
        assert m['prompt'] in prompt
        assert '위원3' not in prompt
        if '이어서 위원1 님이' in prompt:
            return '예산 심사를 시작하겠습니다. 검토 자료 확인이 필요합니다.'
        assert '예산 심사를 시작하겠습니다.' in prompt
        return '검토 자료가 있습니까? 안내 감사합니다.'
    result = generate_local(m, tmp_path, loader=lambda _: (None, None), completion=complete)
    assert '검토 자료가 있습니까?' in result
    assert '위원3' not in result
    assert m == original


@pytest.mark.parametrize('content', ['', '미등록인: 제가 승인합니다.', '{"speeches": []}'])
def test_prompt_rejects_invalid_turn(meeting, content):
    from business_planning.local_ai import generate_local
    from business_planning.minutes import prompt_meeting
    m = prompt_meeting(meeting)
    original = copy.deepcopy(m)
    with pytest.raises(ValueError, match='발언 형식'):
        generate_local(m, loader=lambda _: (None, None), completion=lambda *_: content)
    assert m == original


def test_prompt_docx_and_migration_preserve_old_data(tmp_path, meeting):
    from docx import Document
    from business_planning.minutes import prompt_meeting, document_text
    original = copy.deepcopy(meeting)
    m = prompt_meeting(meeting)
    assert meeting == original and m['legacy_source'] == original
    m.update(quarter=4, final='위원1: 확인하겠습니다.\n위원2: 감사합니다.', confirmed=True)
    path = tmp_path / 'prompt.docx'
    export_docx(m, path)
    text = '\n'.join(p.text for p in Document(path).paragraphs)
    assert '4분기' in text and m['final'] in text
    assert '개회선언' not in text and '검토 중입니다.' not in text
    assert m['final'] in document_text(m)
    store = MinutesStore(tmp_path)
    store.save(m)
    assert store.load(m['id'])['legacy_source'] == original


def test_prompt_accepts_single_speech_envelope_and_assigns_registered_speakers(meeting):
    from business_planning.local_ai import generate_local
    from business_planning.minutes import prompt_meeting
    m = prompt_meeting(meeting)
    count = []
    def completion(*args):
        person = '위원1' if len(count) % 2 == 0 else '위원2'
        count.append(person)
        return '```json\n' + json.dumps({'name': person, 'content': '자료를 확인하겠습니다.'}) + '\n```'
    text = generate_local(m, loader=lambda _: (None, None), completion=completion)
    assert count == ['위원1', '위원2'] * 3
    assert text.count('위원1(시설장):') == 3
    assert text.count('위원2(대표):') == 3


def test_prompt_requires_input_before_model_load(meeting):
    from business_planning.local_ai import generate_local
    from business_planning.minutes import prompt_meeting
    m = prompt_meeting(meeting)
    m['prompt'] = ' '
    def no_load(_):
        pytest.fail('빈 프롬프트에서는 모델을 로딩하지 않아야 합니다.')
    with pytest.raises(ValueError, match='프롬프트'):
        generate_local(m, loader=no_load)
