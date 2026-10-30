from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook
from selenium.common.exceptions import NoSuchWindowException

from ui.pages import federation_tool as module
from ui.pages.macro_table_extract import SCROLL_SCRIPT, TABLE_SCRIPT, TableWorkbook, collect_scrolled_table


def test_collector_retains_merged_rows_before_browser_closes():
    first = {'header': True, 'rows': [['name'], ['A'], ['B']],
             'scroll': {'top': True, 'bottom': False}}
    second = {'header': True, 'rows': [['name'], ['B'], ['C']],
              'scroll': {'top': False, 'bottom': False}}
    snapshots = []
    with pytest.raises(NoSuchWindowException):
        collect_scrolled_table(lambda: first, Mock(),
                               Mock(side_effect=[second, NoSuchWindowException()]),
                               Mock(), on_progress=snapshots.append)
    assert snapshots[-1]['rows'] == [['name'], ['A'], ['B'], ['C']]


@pytest.mark.parametrize('batch', [False, True])
def test_browser_close_saves_current_rows_and_preserves_prior_branches(tmp_path, monkeypatch, batch):
    owner = module.InvoiceProcessor.__new__(module.InvoiceProcessor)
    owner.driver = Mock()
    owner.config = {}
    owner.timeout = 1
    owner._log = Mock()
    owner.check_branch_stop_alert = Mock()
    owner.export_branch_name = 'B'
    workbook = TableWorkbook(tmp_path, 'partial')
    if batch:
        workbook.add_branch({'rows': [['old']]}, 'staff', 'staff', 'A')
        owner.export_session = {'workbook': workbook}
    monkeypatch.setattr(module, 'TableWorkbook', lambda *args: workbook)
    opened = Mock()
    monkeypatch.setattr(module.os, 'startfile', opened)
    element = Mock()
    owner.driver.find_element.return_value = element

    def execute(script, *args):
        if script == TABLE_SCRIPT:
            return {'rows': [['new']]}
        if script == SCROLL_SCRIPT:
            if args[-1] == 0:
                return {'kind': 'native', 'top': True, 'bottom': False}
            raise NoSuchWindowException('user closed browser')
        return element

    owner.driver.execute_script.side_effect = execute

    class Wait:
        def __init__(self, *args):
            pass

        def until(self, predicate):
            return predicate(owner.driver)

    monkeypatch.setattr(module, 'TimedWebDriverWait', Wait)
    selector = {'by': 'xpath', 'value': '//grid', 'action': 'save_xlsx', 'output_name': 'staff'}
    owner.run_workflow_items_once = lambda items: owner.extract_table('staff', selector)
    with pytest.raises(NoSuchWindowException):
        owner.run_configured_workflow()
    saved = load_workbook(workbook.path)
    try:
        rows = list(saved.active.values)
        if batch:
            assert rows[1:] == [('A', 'old'), ('B', 'new')]
            assert owner.export_session['open_requested'] is True
            opened.assert_not_called()
        else:
            assert rows[1:] == [('new',)]
            opened.assert_called_once_with(str(workbook.path))
    finally:
        saved.close()


def test_decision_menu_uses_scoped_text_and_pointer():
    import json
    path = Path(__file__).resolve().parents[1] / 'data/federation_selectors_adjustment_analysis.json'
    step = json.loads(path.read_text(encoding='utf-8'))['selectors']['element_013']
    assert '.grid_menu.body.' in step['value']
    assert "normalize-space(.)='결정내용조회'" in step['value']
    assert step['action'] == 'pointer_click'
