import time

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication, QComboBox, QDialogButtonBox, QFileDialog, QMessageBox, QPushButton, QTableWidget, QTabWidget
from PySide6.QtCore import QMimeData, QPoint, QPointF, Qt, QTimer, QUrl
from PySide6.QtGui import QDragEnterEvent, QDropEvent

from ui.pages.excel_data_sort_page import ExcelDataSortPage


@pytest.fixture(scope="module")
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def page(app, tmp_path, monkeypatch):
    def fail_message(_parent, _title, message, *_args):
        pytest.fail(message)

    monkeypatch.setattr(QMessageBox, "warning", fail_message)
    widget = ExcelDataSortPage(tmp_path / "templates")
    yield widget
    wait_worker(app, widget)
    widget.close()
    widget.deleteLater()
    app.processEvents()


def wait_worker(app, widget):
    deadline = time.monotonic() + 15
    while (widget.worker is not None or widget.preview_timer.isActive()) and time.monotonic() < deadline:
        app.processEvents()
        time.sleep(0.005)
    assert widget.worker is None, "Excel worker timed out"
    assert not widget.preview_timer.isActive(), "Automatic preview timed out"


def create_source(path, value):
    wb = Workbook()
    wb.active.title = "급여"
    wb.active.append(["직원명", "금액"])
    wb.active.append([value, 100])
    wb.save(path)
    wb.close()


def configure(page):
    page.name_edit.setText("노인급여")
    page.read_columns()
    page.columns_table.item(0, 3).setText("이름")


def test_ui_save_recognize_multifile_preview_export(page, app, tmp_path, monkeypatch):
    first, second = tmp_path / "첫째.xlsx", tmp_path / "둘째.xlsx"
    create_source(first, "가")
    create_source(second, "나")
    page.load_paths([first])
    wait_worker(app, page)
    assert page.entries[0]["template"] is None
    configure(page)
    wait_worker(app, page)
    assert page.dataset_table.model().rowCount() == 1
    page.save_template()
    assert (tmp_path / "templates" / "노인급여.json").exists()
    page.load_paths([first, second])
    wait_worker(app, page)
    assert len(page.entries) == 2
    assert page.entries[1]["template"].name == "노인급여"
    assert "자동 선택" in page.entries[1]["status"]
    page.files.setCurrentRow(1)
    page.preview()
    wait_worker(app, page)
    assert page.dataset_table.model().rowCount() == 1
    assert page.result.datasets[0].rows[0][:2] == ["나", 100]
    destination = tmp_path / "Dataset.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_args: (str(destination), ""))
    page.export_all()
    wait_worker(app, page)
    wb = load_workbook(destination)
    try:
        assert list(wb.active.values)[0][:2] == ("이름", "금액")
        assert [row[:2] for row in list(wb.active.values)[1:]] == [("가", 100), ("나", 100)]
    finally:
        wb.close()


def test_file_switch_retains_unsaved_settings_and_manual_template_selection(page, app, tmp_path):
    first, second = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    create_source(first, "가")
    create_source(second, "나")
    page.load_paths([first, second])
    wait_worker(app, page)
    configure(page)
    page.files.setCurrentRow(1)
    assert page.entries[1]["template"] is None
    page.files.setCurrentRow(0)
    assert page.name_edit.text() == "노인급여"
    assert page.columns_table.item(0, 3).text() == "이름"
    page.save_template()
    page.files.setCurrentRow(1)
    page.template_list.setCurrentRow(0)
    page.apply_template()
    assert page.entries[1]["template"].name == "노인급여"
    assert "수동 선택" in page.entries[1]["status"]


def test_header_column_range_for_horizontal_repeated_tables(page, app, tmp_path):
    path = tmp_path / "가로.xlsx"
    wb = Workbook()
    wb.active.append(["직원명", "금액", None, "직원명", "금액"])
    wb.active.append(["가", 100, None, "나", 200])
    wb.save(path)
    wb.close()
    page.load_paths([path])
    wait_worker(app, page)
    page.last_column_spin.setValue(2)
    configure(page)
    assert page.columns_table.rowCount() == 2
    page.preview()
    wait_worker(app, page)
    assert [row[:2] for row in page.result.datasets[0].rows] == [["가", 100], ["나", 200]]


def test_drop_files_and_ignore_unsupported_and_duplicates(page, app, tmp_path):
    path = tmp_path / "드롭.XLSX"
    create_source(path, "가")
    other = tmp_path / "note.txt"
    other.write_text("test", encoding="utf-8")
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(path)), QUrl.fromLocalFile(str(other)),
                  QUrl("https://example.com/a.xlsx")])
    enter = QDragEnterEvent(QPoint(5, 5), Qt.CopyAction, mime, Qt.LeftButton, Qt.NoModifier)
    QApplication.sendEvent(page.files.viewport(), enter)
    assert enter.isAccepted()
    event = QDropEvent(QPointF(5, 5), Qt.CopyAction, mime, Qt.LeftButton, Qt.NoModifier)
    QApplication.sendEvent(page.files.viewport(), event)
    assert event.isAccepted()
    wait_worker(app, page)
    assert len(page.entries) == 1
    page.files.files_dropped.emit([str(path)])
    wait_worker(app, page)
    assert len(page.entries) == 1
    mime.setUrls([QUrl.fromLocalFile(str(other))])
    enter = QDragEnterEvent(QPoint(5, 5), Qt.CopyAction, mime, Qt.LeftButton, Qt.NoModifier)
    QApplication.sendEvent(page.files.viewport(), enter)
    assert not enter.isAccepted()


def test_added_column_ui_insert_save_switch_preview_export(page, app, tmp_path, monkeypatch):
    paths = [tmp_path / "first.xlsx", tmp_path / "second.xlsx"]
    for path in paths:
        create_source(path, "가")
    page.load_paths(paths)
    wait_worker(app, page)
    configure(page)
    page.added_preview.selectColumn(0)
    page.add_column()
    page.added_table.item(0, 1).setText("연월")
    page.added_table.item(0, 3).setText("2026-09")
    page.added_preview.selectColumn(1)
    page.add_column(after=True)
    page.added_table.item(1, 1).setText("참조")
    page.added_table.cellWidget(1, 2).setCurrentIndex(1)
    page.added_table.item(1, 5).setText("B2")
    assert [column.position for column in page.editor_template().added_columns] == [1, 3]
    page.files.setCurrentRow(1)
    assert page.added_table.rowCount() == 0
    page.files.setCurrentRow(0)
    assert page.added_table.item(0, 3).text() == "2026-09"
    page.save_template()
    assert len(page.store.load_all()[0][0].added_columns) == 2
    page.preview()
    wait_worker(app, page)
    assert page.result.datasets[0].rows[0][:4] == ["2026-09", "가", 100, 100]
    page.template_list.setCurrentRow(0)
    page.apply_template_all()
    destination = tmp_path / "ui_added.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_args: (str(destination), ""))
    page.export_all()
    wait_worker(app, page)
    wb = load_workbook(destination)
    try:
        assert list(wb.active.values)[0][:4] == ("연월", "이름", "참조", "금액")
        assert len(list(wb.active.values)) == 3
    finally:
        wb.close()
    page.added_table.selectRow(0)
    page.remove_added_column()
    assert page.editor_template().added_columns[0].position == 2


def test_mapping_dialog_recognition_preview_and_save(page, app, tmp_path, monkeypatch):
    from excel_data_sort import branch_lookup

    directory = tmp_path / "branches"
    directory.mkdir()
    (directory / "branches.json").write_text('[{"organization_code": "100", "branch_name": "Branch", "corporation_name": "Company"}]', encoding="utf-8")
    monkeypatch.setattr(branch_lookup, "DATA_DIRECTORY", directory)
    path = tmp_path / "mapping.xlsx"
    create_source(path, "가")
    page.load_paths([path])
    wait_worker(app, page)
    configure(page)
    page.add_column()
    page.added_table.item(0, 1).setText("지점명")
    page.added_table.item(0, 5).setText("B2")
    page.added_table.cellWidget(0, 2).setCurrentIndex(1)
    assert page.added_table.item(0, 6).text() == "100"

    def edit_dialog():
        dialog = app.activeModalWidget()
        combo = dialog.findChild(QComboBox, "branch_return_field")
        combo.setCurrentIndex(combo.findData("corporation_name"))
        dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Save).click()

    QTimer.singleShot(0, edit_dialog)
    page.edit_value_mapping(0)
    assert page.added_table.item(0, 6).text() == "100"
    assert page.added_table.item(0, 7).text() == "Company"
    page.save_template()
    assert page.store.load_all()[0][0].added_columns[0].branch_lookup == {"match_field": "organization_code", "return_field": "corporation_name"}
    page.preview()
    wait_worker(app, page)
    assert page.result.datasets[0].rows[0][0] == "Company"


def test_dataset_auto_refresh_and_consolidated_controls(page, app, tmp_path):
    path = tmp_path / "auto.xlsx"
    create_source(path, "가")
    page.load_paths([path])
    wait_worker(app, page)
    configure(page)
    wait_worker(app, page)
    assert page.result.datasets[0].columns[0] == "이름"
    page.columns_table.item(0, 3).setText("변경 이름")
    wait_worker(app, page)
    assert page.result.datasets[0].columns[0] == "변경 이름"
    page.columns_table.item(0, 3).setText("")
    wait_worker(app, page)
    assert page.result is None
    assert "출력 컬럼명" in page.dataset_hint.text()
    page.columns_table.item(0, 3).setText("이름")
    wait_worker(app, page)
    assert page.result is not None
    labels = [button.text() for button in page.findChildren(QPushButton)]
    assert labels.count("양식 JSON 저장") == 1
    assert "변환 미리보기" not in labels and "Normalize 미리보기" not in labels
    assert all(tabs.tabText(i) != "Normalize 결과" for tabs in page.findChildren(QTabWidget) for i in range(tabs.count()))
