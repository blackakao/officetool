import json
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from ui.pages import federation_tool as module
from ui.pages.macro_table_extract import (
    SCROLL_SCRIPT, TABLE_SCRIPT, TableWorkbook, collect_scrolled_table, merge_table_rows,
    save_extracted_table, table_wheel_events,
)


def test_xlsx_preserves_blank_columns_dates_and_formula_like_text(tmp_path):
    result = {"header": True, "rows": [
        ["NO", "성명", "생년월일", "사특 가산 제외", "입력여부"],
        ["001", "테스트", "76-08-15", "", "O"],
        ["002", "=1+1", "62-05-19", "", "X"],
    ]}
    path, columns, rows = save_extracted_table(result, tmp_path, "종사자 입사보고")
    book = load_workbook(path)
    try:
        sheet = book.active
        assert list(sheet.values) == [tuple(columns), ("001", "테스트", "76-08-15", None, "O"),
                                      ("002", "=1+1", "62-05-19", None, "X")]
        assert sheet["B3"].data_type == "s"
        assert sheet["A2"].data_type == "s"
    finally:
        book.close()
    second, _, _ = save_extracted_table(result, tmp_path, "종사자 입사보고")
    assert second != path and path.exists()
    assert rows[0][3] == ""


def test_body_only_table_does_not_discard_first_record(tmp_path):
    _, columns, rows = save_extracted_table({"rows": [["1", "첫 행"], ["2"]]}, tmp_path, "body")
    assert columns == ["열 1", "열 2"]
    assert rows == [["1", "첫 행"], ["2", ""]]
    with pytest.raises(ValueError, match="텍스트"):
        save_extracted_table({"rows": [["", ""]]}, tmp_path, "empty")


def test_extract_queues_file_until_workflow_finishes(tmp_path, monkeypatch):
    owner = module.InvoiceProcessor.__new__(module.InvoiceProcessor)
    owner.driver = Mock()
    owner.config = {}
    owner.timeout = 20
    owner._log = Mock()
    element = Mock()
    element.is_displayed.return_value = True
    owner.driver.find_element.return_value = element
    def execute(script, *args):
        if script == TABLE_SCRIPT:
            return {"rows": [["1", "테스트", "76-08-15", "", "O"]]}
        if script == SCROLL_SCRIPT:
            return {"kind": "none", "top": True, "bottom": True}
        return element
    owner.driver.execute_script.side_effect = execute
    monkeypatch.setattr(module, "TableWorkbook", lambda directory, name: TableWorkbook(tmp_path, name))
    opened = Mock()
    monkeypatch.setattr(module.os, "startfile", opened)

    class ImmediateWait:
        def __init__(self, *args):
            pass

        def until(self, predicate):
            return predicate(owner.driver)

    monkeypatch.setattr(module, "TimedWebDriverWait", ImmediateWait)
    path = owner._run_selector_step("extract", {"type": "table_extract", "by": "xpath", "value": "//grid", "action": "save_xlsx_open"})
    assert path.exists()
    assert owner.extracted_tables["extract"]["rows"][0][3] == ""
    second = owner._run_selector_step("extract2", {"type": "table_extract", "by": "xpath", "value": "//grid", "action": "save_xlsx_open"})
    assert second == path
    assert len(list(tmp_path.glob('*.xlsx'))) == 1
    book = load_workbook(path)
    assert len(book.sheetnames) == 2
    book.close()
    opened.assert_not_called()
    owner.run_workflow_items_once = Mock(side_effect=ValueError("later step failed"))
    with pytest.raises(ValueError):
        owner.run_configured_workflow()
    opened.assert_not_called()
    owner.run_workflow_items_once = Mock()
    owner.run_configured_workflow()
    opened.assert_called_once_with(str(path))
    owner.open_extracted_files()
    opened.assert_called_once()


def test_extract_editor_roundtrip_and_adjustment_step(tmp_path):
    from PySide6.QtWidgets import QApplication
    from ui.pages.federation_config import EMPTY_SELECTOR_TEMPLATE
    from ui.pages.federation_dialogs import SelectorConfigDialog

    app = QApplication.instance() or QApplication([])
    path = Path(__file__).resolve().parents[1] / "data" / "federation_selectors_adjustment_analysis.json"
    selectors = list(json.loads(path.read_text(encoding="utf-8"))["selectors"].values())
    index = next(i for i, step in enumerate(selectors) if step["type"] == "table_extract")
    step = selectors[index]
    assert selectors[index - 1]["label"] == "근무내용관리"
    assert selectors[index + 1]["label"] == "근무내용제출"
    assert step["action"] == "save_xlsx_open"
    dialog = SelectorConfigDialog(None, tmp_path / "edit.json", EMPTY_SELECTOR_TEMPLATE)
    dialog.table.setRowCount(1)
    dialog._set_row(0, "extract", step)
    saved = dialog._row_data(0)[1]
    for key in ("type", "by", "value", "output_name", "action", "scroll", "max_scrolls"):
        assert saved[key] == step[key]
    assert dialog.table.cellWidget(0, 2).isEnabled()
    dialog.close()
    assert app is not None


def test_virtual_grid_rewinds_collects_97_rows_and_exports_xlsx(tmp_path):
    position = [80]
    source = [[str(i), f"테스트{i}", "76-08-15", "", "O"] for i in range(1, 98)]
    header = ["NO", "성명", "생년월일", "제외", "입력여부"]

    def read():
        return {"header": True, "rows": [header] + source[position[0]:position[0]+16],
                "scroll": {"kind": "wheel", "position": position[0], "top": position[0] == 0,
                           "bottom": position[0] == 81}}

    def move(direction):
        position[0] = max(0, min(81, position[0] + direction * 8))

    result = collect_scrolled_table(read, move, lambda previous: read(), Mock())
    assert result["rows"] == [header] + source
    path, _, rows = save_extracted_table(result, tmp_path, "97행")
    assert len(rows) == 97
    book = load_workbook(path, read_only=True)
    try:
        assert book.active.max_row == 98
        assert book.active["A2"].value == "1"
        assert book.active["A98"].value == "97"
    finally:
        book.close()


def test_overlap_merge_keeps_equal_records_in_different_rows():
    assert merge_table_rows([["a"], ["same"], ["same"], ["b"]],
                            [["same"], ["b"], ["same"], ["c"]]) == [
                                ["a"], ["same"], ["same"], ["b"], ["same"], ["c"]]


def test_native_full_dom_does_not_duplicate_rows_on_scroll():
    position = [0]
    def read():
        return {"rows": [["1"], ["2"]], "scroll": {
            "kind": "native", "position": position[0], "top": position[0] == 0,
            "bottom": position[0] == 2}}
    def move(direction):
        position[0] = max(0, min(2, position[0] + direction))
    assert collect_scrolled_table(read, move, lambda previous: read(), Mock())["rows"] == [["1"], ["2"]]


@pytest.mark.parametrize("max_scrolls", [2, 500])
def test_stalled_grid_or_scroll_limit_is_not_saved_as_complete(max_scrolls):
    read = lambda: {"rows": [["1"]], "scroll": {"top": True, "bottom": False, "position": 0}}
    with pytest.raises(ValueError, match="부분 결과"):
        collect_scrolled_table(read, Mock(), lambda previous: read(), Mock(), max_scrolls=max_scrolls)


def test_adjustment_has_seven_scroll_exports_without_redundant_waits():
    path = Path(__file__).resolve().parents[1] / "data" / "federation_selectors_adjustment_analysis.json"
    steps = list(json.loads(path.read_text(encoding="utf-8"))["selectors"].values())
    exports = [step for step in steps if step['type'] == 'table_extract']
    assert len(exports) == 7
    assert all(step['scroll'] and step['action'] == 'save_xlsx_open' for step in exports)
    assert not any(step.get('action') == 'short_wait' for step in steps)
    for label, count in [('기타자료관리', 1), ('결정내용조회', 2)]:
        index = next(i for i, step in enumerate(steps) if step['label'] == label)
        assert all(step['type'] == 'table_extract' for step in steps[index-count:index])
    assert all(step['type'] == 'table_extract' for step in steps[-5:-2])


def test_workbook_keeps_seven_sheets_and_prior_checkpoint_on_save_failure(tmp_path, monkeypatch):
    book = TableWorkbook(tmp_path, '한 작업')
    for index in range(7):
        book.add({'rows': [['001', str(index), '=1+1', '']]}, '중복/표')
    saved = load_workbook(book.path)
    try:
        assert len(saved.sheetnames) == 7
        assert len(set(saved.sheetnames)) == 7
        for index, sheet in enumerate(saved):
            assert sheet['A2'].value == '001'
            assert sheet['B2'].value == str(index)
            assert sheet['C2'].data_type == 's'
    finally:
        saved.close()
    from ui.pages import macro_table_extract as extract_module
    monkeypatch.setattr(extract_module, 'export_excel', Mock(side_effect=PermissionError('locked')))
    with pytest.raises(PermissionError):
        book.add({'rows': [['8']]}, '실패')
    assert len(book.datasets) == 7
    saved = load_workbook(book.path)
    assert len(saved.sheetnames) == 7
    saved.close()
    assert TableWorkbook(tmp_path, '한 작업').path != book.path


@pytest.mark.parametrize('count,expected', [(4, 1), (12, 2), (16, 3), (30, 6)])
def test_wheel_bursts_leave_rows_for_overlap(count, expected):
    snapshot = {'header': True, 'rows': [['NO']] + [[str(i)] for i in range(count)]}
    assert table_wheel_events(snapshot, 1) == expected
    assert table_wheel_events(snapshot, -1) == 8


def test_branch_rows_append_to_same_table_sheet_and_align_columns(tmp_path):
    book = TableWorkbook(tmp_path, '다중 지점')
    book.add_branch({'header': True, 'rows': [['NO', '성명'], ['001', '같은이름']]}, '입사보고', 'staff', 'A지점')
    book.add_branch({'header': True, 'rows': [['성명', 'NO', '입력여부'], ['같은이름', '001', 'O']]}, '입사보고', 'staff', 'B지점')
    book.add_branch({'rows': [['보험1']]}, '보험', 'insurance', 'A지점')
    book.add_branch({'rows': [['보험2']]}, '보험', 'insurance', 'B지점')
    saved = load_workbook(book.path)
    try:
        assert saved.sheetnames == ['입사보고', '보험']
        assert list(saved['입사보고'].values) == [
            ('지점명', 'NO', '성명', '입력여부'),
            ('A지점', '001', '같은이름', None), ('B지점', '001', '같은이름', 'O')]
        assert list(saved['보험'].values) == [('지점명', '열 1'), ('A지점', '보험1'), ('B지점', '보험2')]
        assert len(list(tmp_path.glob('*.xlsx'))) == 1
    finally:
        saved.close()


def test_branch_append_failure_does_not_mutate_prior_rows(tmp_path, monkeypatch):
    from ui.pages import macro_table_extract as extract_module
    book = TableWorkbook(tmp_path, 'partial')
    result = {'header': True, 'rows': [['값', '값'], ['001', '=1+1']]}
    book.add_branch(result, '표', 'table', 'A지점')
    before = book.path.read_bytes()
    monkeypatch.setattr(extract_module, 'export_excel', Mock(side_effect=PermissionError('locked')))
    with pytest.raises(PermissionError):
        book.add_branch(result, '표', 'table', 'B지점')
    assert book.datasets[0].rows == [['A지점', '001', '=1+1']]
    assert book.path.read_bytes() == before


def test_branch_rows_append_to_same_table_sheet_and_align_columns(tmp_path):
    book = TableWorkbook(tmp_path, '다중 지점')
    book.add_branch({'header': True, 'rows': [['NO', '성명'], ['001', '같은이름']]}, '입사보고', 'staff', 'A지점')
    book.add_branch({'header': True, 'rows': [['성명', 'NO', '입력여부'], ['같은이름', '001', 'O']]}, '입사보고', 'staff', 'B지점')
    book.add_branch({'rows': [['보험1']]}, '보험', 'insurance', 'A지점')
    book.add_branch({'rows': [['보험2']]}, '보험', 'insurance', 'B지점')
    saved = load_workbook(book.path)
    try:
        assert saved.sheetnames == ['입사보고', '보험']
        assert list(saved['입사보고'].values) == [
            ('지점명', 'NO', '성명', '입력여부'),
            ('A지점', '001', '같은이름', None), ('B지점', '001', '같은이름', 'O')]
        assert list(saved['보험'].values) == [('지점명', '열 1'), ('A지점', '보험1'), ('B지점', '보험2')]
        assert len(list(tmp_path.glob('*.xlsx'))) == 1
    finally:
        saved.close()


def test_branch_append_failure_does_not_mutate_prior_rows(tmp_path, monkeypatch):
    from ui.pages import macro_table_extract as extract_module
    book = TableWorkbook(tmp_path, 'partial')
    result = {'header': True, 'rows': [['값', '값'], ['001', '=1+1']]}
    book.add_branch(result, '표', 'table', 'A지점')
    before = book.path.read_bytes()
    monkeypatch.setattr(extract_module, 'export_excel', Mock(side_effect=PermissionError('locked')))
    with pytest.raises(PermissionError):
        book.add_branch(result, '표', 'table', 'B지점')
    assert book.datasets[0].rows == [['A지점', '001', '=1+1']]
    assert book.path.read_bytes() == before


@pytest.mark.parametrize('snapshot', [
    {'rows': []}, {'header': True, 'rows': [['이름']]},
    {'rows': [['', '  ']]},
])
@pytest.mark.parametrize('skip_empty', [False, True])
def test_empty_table_skip_is_opt_in_and_only_after_wait(monkeypatch, snapshot, skip_empty):
    from selenium.common.exceptions import TimeoutException
    owner = module.InvoiceProcessor.__new__(module.InvoiceProcessor)
    owner.driver = Mock()
    owner.timeout = 20
    owner._log = Mock()
    owner.check_branch_stop_alert = Mock()
    root = Mock()
    root.is_displayed.return_value = True
    owner.driver.execute_script.side_effect = [root, snapshot, snapshot]
    calls = []
    class Wait:
        def __init__(self, *args):
            pass
        def until(self, condition):
            calls.append(condition)
            if len(calls) == 1:
                return root
            assert not condition(owner.driver)
            assert not condition(owner.driver)
            raise TimeoutException()
    monkeypatch.setattr(module, 'TimedWebDriverWait', Wait)
    selector = {'by': 'xpath', 'value': '//grid', 'skip_empty': skip_empty}
    if skip_empty:
        assert owner.extract_table('empty', selector) is None
        assert '[table_skip_empty]' in owner._log.call_args[0][0]
        owner.check_branch_stop_alert.assert_called_once()
        assert not hasattr(owner, 'table_workbook')
    else:
        with pytest.raises(TimeoutException):
            owner.extract_table('empty', selector)


def test_missing_table_is_not_skipped(monkeypatch):
    from selenium.common.exceptions import TimeoutException
    owner = module.InvoiceProcessor.__new__(module.InvoiceProcessor)
    owner.timeout = 20
    owner.driver = Mock()
    class Wait:
        def __init__(self, *args):
            pass
        def until(self, condition):
            raise TimeoutException('missing screen')
    monkeypatch.setattr(module, 'TimedWebDriverWait', Wait)
    with pytest.raises(TimeoutException, match='missing screen'):
        owner.extract_table('missing', {'by': 'xpath', 'value': '//grid', 'skip_empty': True})
    owner.driver.execute_script.assert_not_called()


@pytest.mark.parametrize('kind', ['native', 'wheel'])
def test_bottom_rechecks_include_delayed_last_four_rows_in_xlsx(tmp_path, kind):
    source = [[str(i), f'직원{i}'] for i in range(1, 46)]
    header = ['NO', '성명']
    phase = [0]
    bottom_polls = [0]
    def snapshot():
        if phase[0] == 0:
            visible = source[:16]
        elif phase[0] == 1:
            visible = source[12:28]
        elif phase[0] == 2:
            visible = source[24:41]
        else:
            visible = source[30:45]
        return {'rows': [header] + visible, 'header': True,
                'scroll': {'kind': kind, 'position': min(phase[0], 2),
                           'top': phase[0] == 0, 'bottom': phase[0] >= 2}}
    def move(direction):
        if phase[0] < 2:
            phase[0] += 1
        else:
            bottom_polls[0] += 1
            if bottom_polls[0] >= 2:
                phase[0] = 3
    result = collect_scrolled_table(snapshot, move, lambda previous: snapshot(), Mock())
    assert result['rows'] == [header] + source
    assert bottom_polls[0] >= 5
    path, _, rows = save_extracted_table(result, tmp_path, '끝행검증')
    assert len(rows) == 45
    book = load_workbook(path, read_only=True)
    try:
        assert book.active.max_row == 46
        assert book.active['A46'].value == '45'
        assert book.active['B46'].value == '직원45'
    finally:
        book.close()
