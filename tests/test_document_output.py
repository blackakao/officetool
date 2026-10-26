from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from ui.pages import document_tool, hwp_document_tool
from ui.pages.document_tool import ContentControlDialog, UnifiedDocumentTool
from ui.pages.hwp_document_tool import HwpAutomation, HwpContentControlDialog


def word_dialog(tmp_path):
    source = tmp_path / "template.docx"
    document = Document()
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "name")
    properties.append(tag)
    control.append(properties)
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "original"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    control.append(content)
    document.element.body.append(control)
    document.save(source)
    dialog = ContentControlDialog.__new__(ContentControlDialog)
    dialog.file_path = source
    dialog.settings = {"fields": {"name": {"type": "text"}}}
    return dialog


def test_native_word_applies_values_preserves_source_and_existing_output(tmp_path):
    dialog = word_dialog(tmp_path)
    original = dialog.file_path.read_bytes()
    output = dialog._create_output({"name": "new value"}, "native", output_title="template.docx")
    saved = output.read_bytes()
    second = dialog._create_output({"name": "second value"}, "native", output_title="template.docx")
    assert output == tmp_path / "maked" / "template.docx"
    assert second.name == "template(1).docx"
    assert Document(output).element.xpath("//w:sdtContent//w:t")[0].text == "new value"
    assert Document(second).element.xpath("//w:sdtContent//w:t")[0].text == "second value"
    assert output.read_bytes() == saved
    assert dialog.file_path.read_bytes() == original


def test_pdf_word_still_converts_filled_copy(tmp_path, monkeypatch):
    dialog = word_dialog(tmp_path)
    original = dialog.file_path.read_bytes()
    def convert(source, target):
        assert Document(source).element.xpath("//w:sdtContent//w:t")[0].text == "filled"
        Path(target).write_bytes(b"pdf")
    dialog._ensure_docx2pdf = lambda: SimpleNamespace(convert=convert)
    monkeypatch.setattr(document_tool, "log", lambda *args, **kwargs: None)
    output = dialog._create_output({"name": "filled"}, "pdf")
    assert output.suffix == ".pdf"
    assert output.read_bytes() == b"pdf"
    assert dialog.file_path.read_bytes() == original


@pytest.mark.parametrize("extension", [".pdf", ".docx", ".hwp"])
def test_duplicate_decline_does_not_overwrite(tmp_path, monkeypatch, extension):
    dialog = ContentControlDialog.__new__(ContentControlDialog)
    dialog.file_path = tmp_path / "source.docx"
    existing = tmp_path / f"result{extension}"
    existing.write_bytes(b"keep")
    monkeypatch.setattr(document_tool.QMessageBox, "question", lambda *args: document_tool.QMessageBox.No)
    assert dialog._resolve_output_pdf_path(tmp_path, "result", True, extension) is None
    assert existing.read_bytes() == b"keep"


@pytest.mark.parametrize("native", [False, True])
def test_hwp_output_routes_format_and_applies_values(tmp_path, monkeypatch, native):
    source = tmp_path / "source.hwp"
    source.write_bytes(b"original")
    dialog = HwpContentControlDialog.__new__(HwpContentControlDialog)
    dialog.file_path = source
    dialog.settings = {"fields": {}}
    dialog.controls = {"name": {}}
    automation = Mock()
    automation.__enter__ = Mock(return_value=automation)
    automation.__exit__ = Mock(return_value=False)
    monkeypatch.setattr(hwp_document_tool, "HwpAutomation", lambda: automation)
    monkeypatch.setattr(hwp_document_tool, "log", lambda *args, **kwargs: None)
    output = dialog._create_output({"name": "filled"}, "native" if native else "pdf")
    automation.open.assert_called_once_with(source)
    automation.put_text.assert_called_once_with("name", "filled")
    if native:
        automation.save_native.assert_called_once_with(output)
        automation.save_pdf.assert_not_called()
    else:
        automation.save_pdf.assert_called_once_with(output)
        automation.save_native.assert_not_called()
    assert output.parent == tmp_path / "maked"
    assert output.suffix == (".hwp" if native else ".pdf")
    assert source.read_bytes() == b"original"


def test_hwp_native_save_uses_save_as(tmp_path):
    automation = HwpAutomation()
    automation.hwp = Mock()
    output = tmp_path / "new.hwp"
    automation.save_native(output)
    automation.hwp.SaveAs.assert_called_once_with(str(output.resolve()), "HWP", "")
    automation.hwp.SaveAs.return_value = False
    with pytest.raises(RuntimeError):
        automation.save_native(output)


@pytest.mark.parametrize("suffix", [".hwp", ".docx"])
def test_format_choice_supports_native_pdf_and_cancel(monkeypatch, suffix):
    dialog = ContentControlDialog.__new__(ContentControlDialog)
    dialog.file_path = Path("source" + suffix)
    for index, accepted, expected in [(0, True, "pdf"), (1, True, "native"), (1, False, None)]:
        def choose(*args):
            assert suffix in args[3][1]
            return args[3][index], accepted
        monkeypatch.setattr(document_tool.QInputDialog, "getItem", choose)
        assert dialog._choose_output_format() == expected


def test_unified_list_and_settings_remain_separate(tmp_path, monkeypatch):
    tool = UnifiedDocumentTool.__new__(UnifiedDocumentTool)
    tool.document_folder = tmp_path / "document"
    tool.document_folder.mkdir()
    hwp_folder = tmp_path / "document_hwp"
    hwp_folder.mkdir()
    word = tool.document_folder / "same.docx"
    hwp = hwp_folder / "same.hwp"
    word.touch()
    hwp.touch()
    (hwp_folder / "maked").mkdir()
    (hwp_folder / "maked" / "output.hwp").touch()
    assert set(tool._document_files()) == {word, tool.document_folder / hwp.name}
    assert tool._document_label(word) == "[Word] same.docx"
    assert tool._document_label(hwp) == "[한글] same.hwp"
    monkeypatch.setattr(document_tool, "user_data_path", lambda name: tmp_path / name)
    monkeypatch.setattr(document_tool, "read_json_file", lambda path, default: {
        "documents": {word.name: {"fields": {"word": {}}}, hwp.name: {"fields": {"hwp": {}}}},
    } if path.name == "document_field_settings.json" else {"documents": {hwp.name: {"fields": {"native": {}}}}})
    assert "word" in tool._load_document_settings(word)["fields"]
    assert "native" in tool._load_document_settings(hwp)["fields"]


def test_cancel_generation_creates_no_file_or_history(tmp_path, monkeypatch):
    dialog = ContentControlDialog.__new__(ContentControlDialog)
    dialog.file_path = tmp_path / "source.docx"
    dialog.title_edit = SimpleNamespace(text=lambda: "result")
    dialog._validate_generation_inputs = lambda: []
    dialog._generation_values = lambda: {"name": "value"}
    dialog._choose_output_format = lambda: None
    dialog._create_output = Mock()
    dialog._save_recent_generation = Mock()
    monkeypatch.setattr(document_tool, "log", lambda *args, **kwargs: None)
    dialog.create_from_controls()
    dialog._create_output.assert_not_called()
    dialog._save_recent_generation.assert_not_called()


def test_batch_native_word_creates_each_row_without_modifying_source(tmp_path):
    dialog = word_dialog(tmp_path)
    original = dialog.file_path.read_bytes()
    dialog.branches = []
    dialog.field_widgets = {}
    dialog._show_progress = lambda *args: Mock()
    outputs = dialog.create_documents_from_rows([{"name": "first"}, {"name": "second"}], "native")
    assert len(outputs) == 2
    assert [Document(path).element.xpath("//w:sdtContent//w:t")[0].text for path in outputs] == ["first", "second"]
    assert dialog.file_path.read_bytes() == original
