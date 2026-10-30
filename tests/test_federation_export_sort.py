import json

from openpyxl import load_workbook

from excel_data_sort.exporter import export_excel
from excel_data_sort.federation_export import sort_federation_export
from excel_data_sort.models import ColumnRule, Dataset, Template


def test_saved_template_transforms_hours_and_preserves_branches(tmp_path):
    template = Template('공단 근무시간', 'old sheet', include_source=False, columns=[
        ColumnRule(1, '성명', '이름'),
        ColumnRule(2, '총계', '근무시간', value_transform='hours'),
    ])
    directory = tmp_path / 'templates'
    directory.mkdir()
    payload = json.dumps(template.to_dict(), ensure_ascii=False)
    (directory / '공단 근무시간.json').write_text(payload, encoding='utf-8')
    source = tmp_path / 'source.xlsx'
    export_excel([Dataset('근무', ['지점명', '성명', '총계 / 시간'],
                          [['A점', '가', '160시간30분'], ['B점', '나', '80시간0분']]),
                  Dataset('기타', ['원문'], [['보존']])], source)
    original = source.read_bytes()
    logs = []
    output = sort_federation_export(source, {'근무': '공단 근무시간'}, logs.append,
                                    template_dir=directory)
    book = load_workbook(output)
    try:
        assert list(book['공단 근무시간'].values) == [
            ('지점명', '이름', '근무시간'), ('A점', '가', 160.5), ('B점', '나', 80)]
        assert book['원본_기타']['A2'].value == '보존'
    finally:
        book.close()
    assert source.read_bytes() == original
    assert (directory / '공단 근무시간.json').read_text(encoding='utf-8') == payload


def test_failed_conversion_keeps_raw_sheet_with_warning(tmp_path):
    source = tmp_path / 'source.xlsx'
    export_excel([Dataset('근무', ['성명'], [['보존']])], source)
    logs = []
    output = sort_federation_export(source, {'근무': '없는 양식'}, logs.append,
                                    template_dir=tmp_path)
    book = load_workbook(output)
    try:
        assert book['미정렬_근무']['A2'].value == '보존'
    finally:
        book.close()
    assert any('[table_sort_warning]' in line for line in logs)
