from datetime import date, time

from openpyxl import Workbook, load_workbook

from attendance_analysis.label_io import read_excel_label, write_excel_result


def test_read_excel_label_as_table(tmp_path):
    path = tmp_path / "attendance.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["지점", "직원명", "날짜", "출근시간"])
    worksheet.append(["양천점", "김혜진", date(2026, 8, 1), time(9, 0)])
    workbook.save(path)

    assert read_excel_label(path) == {
        "columns": ["지점", "직원명", "날짜", "출근시간"],
        "rows": [{
            "지점": "양천점",
            "직원명": "김혜진",
            "날짜": "2026-08-01",
            "출근시간": "09:00:00",
        }],
    }


def test_write_excel_result(tmp_path):
    path = tmp_path / "result.xlsx"
    result = {
        "columns": ["지점", "직원명"],
        "rows": [{"지점": "양천점", "직원명": "김혜진"}],
    }

    write_excel_result(result, path)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert list(workbook.active.values) == [("지점", "직원명"), ("양천점", "김혜진")]
    finally:
        workbook.close()
