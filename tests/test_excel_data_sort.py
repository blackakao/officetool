from copy import deepcopy
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from excel_data_sort.engine import convert, convert_many
from excel_data_sort.exporter import export_excel
from excel_data_sort.models import AddedColumn, ColumnRule, Dataset, SheetData, Template, TemplateStore, WorkbookData
from excel_data_sort.normalize import normalize
from excel_data_sort.readers import read_workbook
from excel_data_sort.recognition import RuleRecognizer, build_signature


def template(**kwargs):
    options = dict(name="급여", sheet_name="1월", columns=[
        ColumnRule(1, "직원명", "이름"), ColumnRule(2, "급여", "금액"),
    ], include_source=False)
    options.update(kwargs)
    return Template(**options)


def book(tmp_path, rows, **kwargs):
    return WorkbookData(tmp_path / "input.xlsx", [SheetData("1월", rows, **kwargs)])


def test_normalize_preserves_coordinates_and_does_not_mutate_original(tmp_path):
    original = SheetData("1월", [[None, "직원명", None, "급여"], [None, "가", None, 10],
                                   [None, None, None, 20], [None] * 4, [None, "나", None, 30]],
                         merges=[(2, 3, 2, 2)], hidden_rows={5}, hidden_columns={4})
    before = deepcopy(original)
    result = normalize(original, template())
    assert result.rows == [["직원명"], ["가"], ["가"]]
    assert result.row_numbers == [1, 2, 3]
    assert result.column_numbers == [2]
    assert result.blank_rows == {4}
    assert original == before
    anchor = normalize(original, template(merge_mode="anchor", include_hidden_rows=True, include_hidden_columns=True))
    assert anchor.rows == [["직원명", "급여"], ["가", 10], [None, 20], ["나", 30]]


def test_xlsx_reader_tracks_merged_grouped_hidden_columns_and_dates(tmp_path):
    path = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "1월"
    ws.append(["직원명", "급여", "날짜"])
    ws.append(["가", 100, datetime(2026, 9, 1)])
    ws.append([None, 200, None])
    ws.merge_cells("A2:A3")
    ws.column_dimensions.group("B", "C", hidden=True)
    ws.row_dimensions[3].hidden = True
    wb.create_sheet("숨김").sheet_state = "hidden"
    wb.save(path)
    wb.close()
    result = read_workbook(path)
    assert result.sheets[0].merges == [(2, 3, 1, 1)]
    assert result.sheets[0].hidden_columns == {2, 3}
    assert result.sheets[0].hidden_rows == {3}
    assert result.sheets[0].rows[1][2] == datetime(2026, 9, 1)
    assert result.sheets[1].hidden


def test_real_xls_reader_and_conversion(tmp_path):
    # xlwt is only needed to construct this legacy binary fixture.
    import xlwt

    path = tmp_path / "source.xls"
    wb = xlwt.Workbook()
    ws = wb.add_sheet("1월")
    ws.write(0, 0, "직원명")
    ws.write(0, 1, "급여")
    ws.write_merge(1, 2, 0, 0, "가")
    ws.write(1, 1, 100)
    ws.write(2, 1, 200)
    ws.row(2).hidden = True
    ws.col(2).hidden = True
    ws.write(1, 3, datetime(2026, 9, 1), xlwt.easyxf(num_format_str="YYYY-MM-DD"))
    wb.save(str(path))
    result = read_workbook(path)
    assert result.sheets[0].merges == [(2, 3, 1, 1)]
    assert result.sheets[0].hidden_rows == {3}
    assert result.sheets[0].hidden_columns == {3}
    assert result.sheets[0].rows[1][3] == datetime(2026, 9, 1)
    assert convert(result, template()).datasets[0].rows == [["가", 100]]
    assert convert(result, template(include_hidden_rows=True)).datasets[0].rows == [["가", 100], ["가", 200]]


def test_repeated_tables_horizontal_and_vertical_and_reordered_sheets(tmp_path):
    source = book(tmp_path, [
        ["직원명", "급여", None, "직원명", "급여"],
        ["가", 10, None, "나", 20],
        [None] * 5,
        ["직원명", "급여", None, "직원명", "급여"],
        ["다", 30, None, "라", 40],
    ])
    source.sheets += [SheetData("2월", [["제목", None], ["급여", "직원명"], [50, "마"]]),
                      SheetData("안내", [["사용 방법"], ["참고"]])]
    result = convert(source, template(include_source=True))
    assert [row[:2] for row in result.datasets[0].rows] == [["가", 10], ["나", 20], ["다", 30], ["라", 40], ["마", 50]]
    assert result.datasets[0].rows[-1][2:] == ["input.xlsx", "2월", 3, 1]
    assert any("안내" in warning for warning in result.warnings)


def test_optional_column_matching_renaming_and_extra_columns(tmp_path):
    source = book(tmp_path, [["급여", "추가", "직원명"], [0, "x", "가"]])
    rule = template(columns=[ColumnRule(1, "직원명", "이름"), ColumnRule(2, "급여", "금액"),
                             ColumnRule(3, "직종", "직종", False)])
    with pytest.raises(ValueError, match="추출된 데이터"):
        convert(source, rule)
    rule.header_match = "selected"
    result = convert(source, rule)
    assert result.datasets[0].columns == ["이름", "금액"]
    assert result.datasets[0].rows == [["가", 0]]


def test_blank_boundary_footer_and_data_offset(tmp_path):
    source = book(tmp_path, [["보고서", None], ["직원명", "급여"], ["설명", None], ["가", 1],
                             [None, None], ["나", 2], ["합계", 3], ["주의", "후기"]])
    rule = template(header_row=2, data_start_row=4, stop_at_blank=False, stop_values=["합계", "소계", "총계"])
    assert convert(source, rule).datasets[0].rows == [["가", 1], ["나", 2]]
    rule.stop_at_blank = True
    assert convert(source, rule).datasets[0].rows == [["가", 1]]
    rule.stop_at_blank = False
    rule.data_end_row = 4
    assert convert(source, rule).datasets[0].rows == [["가", 1]]


@pytest.mark.parametrize("suffix", [".xlsx", ".xls"])
def test_default_reads_past_resident_18_to_last_value(tmp_path, suffix):
    path = tmp_path / ("residents" + suffix)
    rows = [["직원명", "급여"]] + [[f"입소자 {n}", n] for n in range(1, 19)]
    rows += [[None, None], ["  ", None], ["입소자 19", 0], ["소계", 19],
             [None, None], ["입소자 20", False]]
    if suffix == ".xlsx":
        wb = Workbook()
        wb.active.title = "1월"
        for row in rows:
            wb.active.append(row)
        wb.active.cell(100, 5).number_format = "0.00"
        wb.save(path)
        wb.close()
    else:
        import xlwt

        wb = xlwt.Workbook()
        ws = wb.add_sheet("1월")
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if value is not None:
                    ws.write(r, c, value)
        ws.write(99, 4, "", xlwt.easyxf(num_format_str="0.00"))
        wb.save(str(path))
    result = convert(read_workbook(path), template(include_source=True))
    assert len(result.datasets[0].rows) == 21
    assert [row[:2] for row in result.datasets[0].rows[-3:]] == [
        ["입소자 19", 0], ["소계", 19], ["입소자 20", False],
    ]
    assert result.datasets[0].rows[-1][-2] == len(rows)
    assert result.normalized[0].row_numbers[-1] == len(rows)


def test_added_columns_positions_references_batch_and_export(tmp_path):
    first = book(tmp_path, [["직원명", "급여"], ["가", 10]])
    first.sheets.append(SheetData("정보", [["2026-09", None]], merges=[(1, 1, 1, 2)]))
    second = deepcopy(first)
    second.path = tmp_path / "second.xlsx"
    second.sheets[1].rows[0][0] = "2026-10"
    rule = template(include_source=True, added_columns=[
        AddedColumn("연월", 1, "cell", sheet_name="정보", cell="B1"),
        AddedColumn("구분", 3, value="입소자"),
        AddedColumn("복사", 5, "cell", cell="B2"),
        AddedColumn("빈값", 6, "cell", cell="Z99"),
    ])
    store = TemplateStore(tmp_path / "templates")
    store.save(rule)
    assert store.load_all() == ([rule], [])
    legacy = template().to_dict()
    legacy.pop("added_columns")
    assert Template.from_dict(legacy).added_columns == []
    datasets, _ = convert_many([(first, rule), (second, rule)])
    assert datasets[0].columns[:6] == ["연월", "이름", "구분", "금액", "복사", "빈값"]
    assert [row[:6] for row in datasets[0].rows] == [
        ["2026-09", "가", "입소자", 10, 10, None],
        ["2026-10", "가", "입소자", 10, 10, None],
    ]
    destination = tmp_path / "added.xlsx"
    export_excel(datasets, destination)
    wb = load_workbook(destination)
    try:
        assert list(wb.active.values)[1][:6] == tuple(datasets[0].rows[0][:6])
    finally:
        wb.close()
    assert first.sheets[0].rows == [["직원명", "급여"], ["가", 10]]
    rule.added_columns[0].sheet_name = "없는 시트"
    with pytest.raises(ValueError, match="참조 시트"):
        convert(first, rule)


@pytest.mark.parametrize("additions", [
    [AddedColumn("이름")], [AddedColumn("새 열", 0)], [AddedColumn("새 열", 4)],
    [AddedColumn("가", 1), AddedColumn("나", 1)],
    [AddedColumn("가", mode="cell", cell="XFE1")],
    [AddedColumn("가", mode="cell", cell="A0")],
])
def test_invalid_added_columns(additions):
    with pytest.raises(ValueError):
        template(added_columns=additions).validate()


def test_value_mapping_roundtrip_batch_and_unmatched_types(tmp_path):
    rule = template(added_columns=[AddedColumn("지점명", 2, "cell", cell="B2",
                                            value_mapping={"100": "서울점", "00100": "부산점"})])
    store = TemplateStore(tmp_path / "mapping")
    store.save(rule)
    restored = store.load_all()[0][0]
    assert restored == rule
    for value, expected in [(100, "서울점"), ("00100", "부산점"), (200, 200), (None, None), (False, False)]:
        source = book(tmp_path, [["직원명", "급여"], ["가", value]])
        assert convert(source, restored).datasets[0].rows == [["가", expected, value]]
    legacy = rule.to_dict()
    legacy["added_columns"][0].pop("value_mapping")
    assert Template.from_dict(legacy).added_columns[0].value_mapping == {}


def test_branch_lookup_live_fields_and_errors(tmp_path, monkeypatch):
    import json
    from excel_data_sort import branch_lookup

    directory = tmp_path / "branches"
    directory.mkdir()
    monkeypatch.setattr(branch_lookup, "DATA_DIRECTORY", directory)
    branches = [{"organization_code": "14161000183", "branch_name": "분당점",
                 "corporation_name": "테스트 법인", "custom_fields": {"custom_phone": "031-123"}}]
    path = directory / "branches.json"
    path.write_text(json.dumps(branches), encoding="utf-8")
    (directory / "branch_fields.json").write_text(json.dumps([
        {"key": "custom_phone", "label": "전화번호", "type": "text"},
        {"key": "custom_image", "label": "직인", "type": "image"},
    ]), encoding="utf-8")
    assert "custom_image" not in branch_lookup.load_branch_data()[1]
    rule = template(added_columns=[AddedColumn("조회 결과", 2, "cell", cell="B2", branch_lookup={
        "match_field": "organization_code", "return_field": "corporation_name"})])
    source = book(tmp_path, [["직원명", "급여"], ["가", 14161000183.0]])
    assert convert(source, rule).datasets[0].rows[0][1] == "테스트 법인"
    branches[0]["corporation_name"] = "변경 법인"
    path.write_text(json.dumps(branches), encoding="utf-8")
    assert convert(source, rule).datasets[0].rows[0][1] == "변경 법인"
    rule.added_columns[0].branch_lookup = {"match_field": "custom_phone", "return_field": "branch_name"}
    source.sheets[0].rows[1][1] = "031-123"
    assert convert(source, rule).datasets[0].rows[0][1] == "분당점"
    assert Template.from_dict(rule.to_dict()) == rule
    branches.append(deepcopy(branches[0]))
    path.write_text(json.dumps(branches), encoding="utf-8")
    with pytest.raises(ValueError, match="중복"):
        convert(source, rule)
    source.sheets[0].rows[1][1] = "unknown"
    with pytest.raises(ValueError, match="일치하는 지점이 없습니다"):
        convert(source, rule)


def test_monthly_weekday_headers_and_recognition_diagnostics(tmp_path):
    from excel_data_sort.recognition import recognition_log

    rule = template(columns=[ColumnRule(1, "성명", "이름"), ColumnRule(2, "1\n(토)", "1일"),
                             ColumnRule(3, "2\n(일)", "2일")])
    original = book(tmp_path, [["성명", "1\n(토)", "2\n(일)"], ["가", 10, 20]])
    rule.signature = build_signature(original, rule)
    shifted = book(tmp_path, [["성명", "1\n(수)", "2\n(목)"], ["가", 30, 40]])
    result = RuleRecognizer().recognize(shifted, [rule])
    assert result.selected == rule
    assert convert(shifted, rule).datasets[0].rows == [["가", 30, 40]]
    shifted.sheets[0].rows[0][2] = "다른 헤더"
    result = RuleRecognizer().recognize(shifted, [rule])
    assert result.selected is None
    details = recognition_log(result)
    assert "2/3" in details and "2(요일)" in details
    assert "기준 80%" in details and "헤더 구조 없음" in details
    assert "원본 1행" in details


def test_separate_and_selected_sheet_and_first_table(tmp_path):
    rows = [["직원명", "급여"], ["가", 1], ["직원명", "급여"], ["나", 2]]
    source = book(tmp_path, rows)
    source.sheets.append(SheetData("2월", rows))
    assert len(convert(source, template(table_mode="separate")).datasets) == 4
    assert len(convert(source, template(sheet_mode="separate")).datasets) == 2
    assert convert(source, template(sheet_mode="selected", table_mode="first")).datasets[0].rows == [["가", 1]]


def test_hidden_sheet_and_hidden_rows_do_not_create_blank_boundaries(tmp_path):
    source = book(tmp_path, [["직원명", "급여"], ["가", 1], ["나", 2], ["다", 3]], hidden_rows={3})
    source.sheets.append(SheetData("숨김", [["직원명", "급여"], ["라", 4]], hidden=True))
    assert convert(source, template()).datasets[0].rows == [["가", 1], ["다", 3]]
    assert len(convert(source, template(include_hidden_sheets=True)).datasets[0].rows) == 3


def test_template_roundtrip_and_invalid_json_is_reported(tmp_path):
    store = TemplateStore(tmp_path / "양식")
    rule = template()
    source = book(tmp_path, [["직원명", "급여"], ["가", 1]])
    rule.signature = build_signature(source, rule)
    path = store.save(rule)
    assert "급여" in path.read_text(encoding="utf-8")
    assert store.load_all() == ([rule], [])
    (store.directory / "broken.json").write_text('{"name": 123}', encoding="utf-8")
    assert len(store.load_all()[1]) == 1
    with pytest.raises(ValueError):
        store.save(template(name="../outside"))
    with pytest.raises(ValueError):
        template(columns=[ColumnRule(1, "A", "같음"), ColumnRule(2, "B", "같음")]).validate()


def test_recognition_selects_similar_layout_and_rejects_ambiguity_and_wrong_headers(tmp_path):
    source = book(tmp_path, [["직원명", "급여"], ["가", 1]])
    first = template()
    first.signature = build_signature(source, first)
    shifted = book(tmp_path, [["보고서", None], ["급 여", "직원명"], [1, "가"]])
    other = template(name="출근부", columns=[ColumnRule(1, "직원명", "이름"), ColumnRule(2, "출근시간", "시간")])
    recognizer = RuleRecognizer()
    assert recognizer.recognize(shifted, [first, other]).selected == first
    duplicate = deepcopy(first)
    duplicate.name = "급여 사본"
    assert recognizer.recognize(source, [first, duplicate]).selected is None
    assert recognizer.recognize(source, [other]).selected is None


def test_export_preserves_types_text_formulas_and_original(tmp_path):
    path = tmp_path / "Dataset.xlsx"
    datasets = [Dataset("A/B", ["이름", "금액", "날짜"], [["=SUM(A1:A2)", 0, datetime(2026, 9, 1)]]),
                Dataset("A:B", ["이름"], [["+문자"]])]
    export_excel(datasets, path)
    wb = load_workbook(path)
    try:
        assert wb.sheetnames == ["A_B", "A_B_2"]
        assert wb.worksheets[0]["A2"].data_type == "s"
        assert wb.worksheets[0]["A2"].value == "=SUM(A1:A2)"
        assert wb.worksheets[0]["B2"].value == 0
        assert wb.worksheets[0]["C2"].value == datetime(2026, 9, 1)
    finally:
        wb.close()
    original = path.read_bytes()
    with pytest.raises(ValueError, match="원본"):
        export_excel(datasets, path, [path])
    assert path.read_bytes() == original
    with pytest.raises(ValueError, match="문자열"):
        export_excel([Dataset("bad", ["a"], [["a\x00b"]])], path)
    assert path.read_bytes() == original


def test_batch_rejects_failed_input_and_keeps_different_schemas_separate(tmp_path):
    first = book(tmp_path, [["직원명", "급여"], ["가", 1]])
    other = book(tmp_path, [["직원명", "급여"], ["나", 2]])
    other.path = tmp_path / "other.xlsx"
    datasets, _warnings = convert_many([(first, template()), (other, template())])
    assert datasets[0].rows == [["가", 1], ["나", 2]]
    renamed = template(columns=[ColumnRule(1, "직원명", "성명"), ColumnRule(2, "급여", "금액")])
    assert len(convert_many([(first, template()), (other, renamed)])[0]) == 2
    failed = book(tmp_path, [["안내"]])
    failed.path = tmp_path / "failed.xlsx"
    with pytest.raises(ValueError, match="failed.xlsx"):
        convert_many([(first, template()), (failed, template())])


def test_cli_end_to_end_and_existing_output_protection(tmp_path, monkeypatch, capsys):
    from excel_data_sort.__main__ import main

    path = tmp_path / "input.xlsx"
    wb = Workbook()
    wb.active.title = "1월"
    wb.active.append(["직원명", "급여"])
    wb.active.append(["가", 10])
    wb.save(path)
    wb.close()
    rule_path = TemplateStore(tmp_path / "templates").save(template())
    output = tmp_path / "result.xlsx"
    monkeypatch.setattr("sys.argv", ["excel_data_sort", str(path), "--template", str(rule_path), "--output", str(output)])
    main()
    assert "저장 완료" in capsys.readouterr().out
    result = load_workbook(output)
    try:
        assert list(result.active.values) == [("이름", "금액"), ("가", 10)]
    finally:
        result.close()
    with pytest.raises(SystemExit) as exc:
        main()
    assert exc.value.code == 2
