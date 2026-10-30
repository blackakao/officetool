import pytest

from excel_data_sort.value_transforms import transform_value
from excel_data_sort.engine import convert
from excel_data_sort.models import AddedColumn, ColumnRule, SheetData, Template, TemplateStore, WorkbookData
from excel_data_sort.exporter import export_excel
from openpyxl import load_workbook


@pytest.mark.parametrize("value,expected", [
    ("11.요양보호사 1급", "요양보호사"), ("8.관리인", "관리인"),
    ("11.요양보호사", "요양보호사"), ("간호조무사(2)", "간호조무사"),
    (" 18. 시설 관리인 ", "시설 관리인"), ("RN-2", "RN"), (None, None), ("", ""),
])
def test_letters(value, expected):
    assert transform_value(value, "letters") == expected


@pytest.mark.parametrize("value,expected", [
    ("160시간20분", 160.33), ("160시간40분", 160.67), ("0시간0분", 0),
    ("20분", 0.33), ("160시간", 160), (" 160 시간 20 분 ", 160.33),
    ("1시간90분", 2.5), (160.335, 160.34), (0, 0), (None, None), ("", ""),
])
def test_hours(value, expected):
    assert transform_value(value, "hours") == expected


@pytest.mark.parametrize("value", ["알 수 없음", "160시간20분추가", "-1시간", True, "NaN", float("inf")])
def test_invalid_hours(value):
    with pytest.raises(ValueError):
        transform_value(value, "hours")


def test_transform_roundtrip_and_numeric_excel(tmp_path):
    rule = Template("시간", "Sheet", columns=[
        ColumnRule(1, "직종", "직종", value_transform="letters"),
        ColumnRule(2, "총계", "총계", value_transform="hours"),
    ], added_columns=[AddedColumn("고정시간", 3, value="0시간20분", value_transform="hours")], include_source=False)
    store = TemplateStore(tmp_path / "templates")
    store.save(rule)
    loaded, errors = store.load_all()
    assert not errors and loaded == [rule]
    book = WorkbookData(tmp_path / "input.xlsx", [SheetData("Sheet", [["직종", "총계"], ["11.요양보호사 1급", "160시간20분"]])])
    result = convert(book, loaded[0])
    assert result.datasets[0].rows == [["요양보호사", 160.33, 0.33]]
    path = tmp_path / "result.xlsx"
    export_excel(result.datasets, path)
    wb = load_workbook(path)
    try:
        assert wb.active.cell(2, 2).value == 160.33
        assert wb.active.cell(2, 2).data_type == "n"
    finally:
        wb.close()
