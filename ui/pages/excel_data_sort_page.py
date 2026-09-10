from copy import deepcopy
from pathlib import Path

from PySide6.QtCore import QAbstractTableModel, QModelIndex, QThread, QTimer, Qt, Signal
from PySide6.QtWidgets import (
    QAbstractItemView, QCheckBox, QComboBox, QDialog, QDialogButtonBox, QFormLayout, QHBoxLayout, QHeaderView,
    QLabel, QLineEdit, QListWidget, QMessageBox, QPushButton, QFileDialog,
    QPlainTextEdit, QScrollArea, QSpinBox, QSplitter, QTableView, QTableWidget,
    QTableWidgetItem, QTabWidget, QVBoxLayout, QWidget,
)
from openpyxl.utils import get_column_letter

from excel_data_sort.engine import added_column_raw_value, added_column_value, convert, convert_many
from excel_data_sort.branch_lookup import load_branch_data, lookup_branch
from excel_data_sort.exporter import export_excel
from excel_data_sort.models import AddedColumn, ColumnRule, Template, TemplateStore, is_empty
from excel_data_sort.normalize import normalize
from excel_data_sort.readers import read_workbook
from excel_data_sort.recognition import RuleRecognizer, build_signature, recognition_log


class ExcelFileList(QListWidget):
    files_dropped = Signal(list)

    def __init__(self):
        super().__init__()
        self.setAcceptDrops(True)
        self.setDragDropMode(QAbstractItemView.DropOnly)
        self.setToolTip("XLSX·XLS 파일을 여기에 드래그해서 추가하세요.")

    @staticmethod
    def dropped_paths(mime):
        return [url.toLocalFile() for url in mime.urls()
                if url.isLocalFile() and Path(url.toLocalFile()).suffix.lower() in {".xlsx", ".xls"}
                and Path(url.toLocalFile()).is_file()]

    def dragEnterEvent(self, event):
        if self.dropped_paths(event.mimeData()):
            event.acceptProposedAction()
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        self.dragEnterEvent(event)

    def dropEvent(self, event):
        paths = self.dropped_paths(event.mimeData())
        if paths:
            event.acceptProposedAction()
            self.files_dropped.emit(paths)
        else:
            event.ignore()


class PreviewModel(QAbstractTableModel):
    def __init__(self, rows=(), columns=(), row_numbers=None, parent=None):
        super().__init__(parent)
        self.rows = rows[:500]
        self.columns = columns[:200]
        self.row_numbers = row_numbers

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.rows)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.columns)

    def data(self, index, role=Qt.DisplayRole):
        if not index.isValid():
            return None
        if role == Qt.TextAlignmentRole:
            return Qt.AlignCenter
        if role != Qt.DisplayRole:
            return None
        value = self.rows[index.row()][index.column()]
        return "" if value is None else str(value)

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role == Qt.DisplayRole:
            if orientation == Qt.Horizontal:
                return self.columns[section]
            return str(self.row_numbers[section] if self.row_numbers else section + 1)
        return None


class ExcelWorker(QThread):
    succeeded = Signal(object)
    failed = Signal(str)

    def __init__(self, action, parent=None):
        super().__init__(parent)
        self.action = action

    def run(self):
        try:
            self.succeeded.emit(self.action())
        except Exception as exc:
            self.failed.emit(str(exc))


class ExcelDataSortPage(QWidget):
    def __init__(self, template_directory=None):
        super().__init__()
        root = Path(__file__).resolve().parents[2]
        self.store = TemplateStore(template_directory or root / "data" / "excel_templates")
        self.recognizer = RuleRecognizer()
        self.entries = []
        self.templates = []
        self.current_index = -1
        self.result = None
        self.worker = None
        self.loading = False
        self.preview_timer = QTimer(self)
        self.preview_timer.setSingleShot(True)
        self.preview_timer.setInterval(350)
        self.preview_timer.timeout.connect(self._auto_preview)
        self.last_preview_error = ""
        layout = QVBoxLayout(self)
        layout.addWidget(QLabel("<h2>엑셀 데이터 정렬</h2>"))
        self.status = QLabel("파일을 추가하고 기준 시트와 헤더를 설정해 주세요.")
        self.status.setWordWrap(True)
        layout.addWidget(self.status)
        self.body = QWidget()
        body_layout = QVBoxLayout(self.body)
        body_layout.setContentsMargins(0, 0, 0, 0)
        self._buttons(body_layout, [("양식 JSON 저장", self.save_template)])
        layout.addWidget(self.body, 1)
        vertical = QSplitter(Qt.Vertical)
        body_layout.addWidget(vertical, 1)
        top = QSplitter(Qt.Horizontal)
        vertical.addWidget(top)
        left = QWidget()
        left_layout = QVBoxLayout(left)
        left_layout.addWidget(QLabel("업로드 파일 (XLSX·XLS 드래그 앤 드롭)"))
        self.files = ExcelFileList()
        self.files.files_dropped.connect(self.load_paths)
        left_layout.addWidget(self.files)
        self._buttons(left_layout, [("파일 추가", self.add_files), ("선택 제거", self.remove_file)])
        left_layout.addWidget(QLabel("등록된 양식"))
        self.template_list = QListWidget()
        left_layout.addWidget(self.template_list)
        self._buttons(left_layout, [("현재 파일에 적용", self.apply_template), ("새 양식", self.new_template)])
        self._buttons(left_layout, [("선택 양식 일괄 적용", self.apply_template_all), ("자동 인식 다시", self.recognize_all)])
        self.template_list.itemDoubleClicked.connect(lambda _item: self.apply_template())
        top.addWidget(left)
        right = QTabWidget()
        top.addWidget(right)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        editor = QWidget()
        editor_layout = QVBoxLayout(editor)
        form = QFormLayout()
        self.name_edit = QLineEdit()
        self.sheet_combo = QComboBox()
        self.header_spin = self._spin(1, 1_048_575, 1)
        self.start_spin = self._spin(2, 1_048_576, 2)
        self.end_spin = self._spin(0, 1_048_576, 0)
        self.end_spin.setSpecialValueText("자동 (마지막 값까지)")
        self.first_column_spin = self._spin(1, 16_384, 1)
        self.last_column_spin = self._spin(0, 16_384, 0)
        self.last_column_spin.setSpecialValueText("마지막 열까지")
        self.merge_combo = self._combo([("병합 영역에 대표값 채우기", "fill"), ("왼쪽 위 값만 유지", "anchor")])
        self.table_combo = self._combo([("같은 헤더의 반복 테이블 합치기", "merge"), ("테이블별 분리", "separate"), ("첫 테이블만", "first")])
        self.sheet_mode_combo = self._combo([("같은 헤더의 시트 합치기", "matching"), ("일치하는 시트를 각각 분리", "separate"), ("기준 시트명만 사용", "selected")])
        self.match_combo = self._combo([("등록된 모든 헤더 일치", "all"), ("사용할 컬럼의 헤더만 일치", "selected")])
        self.hidden_rows = QCheckBox("숨김 행 포함")
        self.hidden_columns = QCheckBox("숨김 열 포함")
        self.hidden_sheets = QCheckBox("숨김 시트 포함")
        self.blank_stop = QCheckBox("빈 행에서 테이블 종료 (해제하면 빈 행을 건너뜀)")
        self.blank_stop.setChecked(False)
        self.stop_edit = QLineEdit()
        self.stop_edit.setPlaceholderText("기본: 없음 (필요 시 합계, 소계, 총계 등 입력)")
        self.source_check = QCheckBox("원본 파일·시트·행·테이블 번호 포함")
        self.source_check.setChecked(True)
        for label, widget in (
            ("양식 이름", self.name_edit), ("기준 시트", self.sheet_combo),
            ("헤더행 (원본 행 번호)", self.header_spin), ("데이터 시작행", self.start_spin),
            ("데이터 끝행 (0 = 자동)", self.end_spin), ("병합 해제 후 값", self.merge_combo),
            ("기준 테이블 첫 열 (A=1)", self.first_column_spin),
            ("기준 테이블 끝 열 (0 = 전체)", self.last_column_spin),
            ("테이블 합치기", self.table_combo), ("시트 합치기", self.sheet_mode_combo),
            ("합치기 판단 기준", self.match_combo), ("", self.hidden_rows),
            ("", self.hidden_columns), ("", self.hidden_sheets), ("", self.blank_stop),
            ("종료 값 (쉼표 구분)", self.stop_edit), ("", self.source_check),
        ):
            form.addRow(label, widget)
        editor_layout.addLayout(form)
        coordinate_help = QLabel(
            "기본 설정은 중간 빈 행을 건너뛰고 값이 있는 마지막 행까지 읽습니다.\n"
            "행 번호는 원본 Excel 기준입니다. 반복 테이블의 시작·끝행은 발견한 헤더와의 간격으로 적용합니다.\n"
            "가로로 반복된 테이블은 기준 테이블 하나의 열 범위를 지정한 뒤 컬럼을 불러오세요.\n"
            "종료 값은 일치하는 테이블 컬럼에서 공백·대소문자를 무시하고 완전히 같은 값일 때 적용됩니다."
        )
        coordinate_help.setWordWrap(True)
        editor_layout.addWidget(coordinate_help)
        self._buttons(editor_layout, [("헤더에서 컬럼 불러오기", self.read_columns)])
        self.columns_table = QTableWidget(0, 4)
        self.columns_table.setHorizontalHeaderLabels(["사용", "원본 열", "헤더명", "출력 컬럼명"])
        self.columns_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.columns_table.setMinimumHeight(190)
        editor_layout.addWidget(self.columns_table)
        scroll.setWidget(editor)
        right.addTab(scroll, "양식 편집 / 설정")
        self.raw_table = self._view()
        right.addTab(self.raw_table, "원본 미리보기")
        self._build_added_tab(right)
        bottom = QTabWidget()
        vertical.addWidget(bottom)
        self.dataset_combo, self.dataset_table = self._preview_tab(bottom, "Dataset 결과")
        self.dataset_hint = QLabel("파일과 헤더를 설정하면 Dataset을 자동으로 갱신합니다.")
        bottom.widget(0).layout().insertWidget(0, self.dataset_hint)
        self.log_view = QPlainTextEdit()
        self.log_view.setReadOnly(True)
        bottom.addTab(self.log_view, "처리 내역")
        body_layout.addWidget(QLabel("미리보기는 최대 500행·200열입니다. Excel 출력에는 전체 Dataset을 저장합니다."))
        self._buttons(body_layout, [("전체 파일 → Excel 저장", self.export_all)])
        top.setSizes([260, 720])
        vertical.setSizes([600, 280])
        self.files.currentRowChanged.connect(self._file_changed)
        self.sheet_combo.currentIndexChanged.connect(self.show_raw)
        self.dataset_combo.currentIndexChanged.connect(self.show_dataset)
        self.header_spin.valueChanged.connect(self._header_changed)
        for widget in (self.name_edit, self.stop_edit):
            widget.textChanged.connect(self.invalidate)
        for widget in (self.sheet_combo, self.merge_combo, self.table_combo, self.sheet_mode_combo, self.match_combo):
            widget.currentIndexChanged.connect(self.invalidate)
        for widget in (self.header_spin, self.start_spin, self.end_spin, self.first_column_spin, self.last_column_spin):
            widget.valueChanged.connect(self.invalidate)
        for widget in (self.hidden_rows, self.hidden_columns, self.hidden_sheets, self.blank_stop, self.source_check):
            widget.toggled.connect(self.invalidate)
        self.columns_table.itemChanged.connect(self.invalidate)
        self.refresh_templates()

    @staticmethod
    def _buttons(layout, actions):
        row = QHBoxLayout()
        for title, action in actions:
            button = QPushButton(title)
            button.clicked.connect(action)
            row.addWidget(button)
        layout.addLayout(row)

    def _build_added_tab(self, tabs):
        widget = QWidget()
        layout = QVBoxLayout(widget)
        help_label = QLabel(
            "아래 미리보기에서 열을 선택한 뒤 앞/뒤에 추가하세요. 위치는 최종 출력 순서(A=1)입니다.\n"
            "고정값은 모든 행에 동일하게 들어갑니다. 셀 참조는 각 파일의 원본 주소(예: B2)를 읽습니다.\n"
            "참조 시트를 비워두면 처리 중인 시트를 사용합니다. 병합 셀은 대표값, 빈 셀은 빈값을 가져옵니다."
        )
        help_label.setWordWrap(True)
        layout.addWidget(help_label)
        self.added_preview = self._view()
        self.added_preview.setSelectionBehavior(QAbstractItemView.SelectColumns)
        self.added_preview.horizontalHeader().sectionClicked.connect(self.added_preview.selectColumn)
        layout.addWidget(self.added_preview, 1)
        self._buttons(layout, [("선택 열 앞에 추가", lambda: self.add_column(False)),
                               ("선택 열 뒤에 추가", lambda: self.add_column(True)),
                               ("선택한 추가 컬럼 삭제", self.remove_added_column)])
        self.added_table = QTableWidget(0, 9)
        self.added_table.setHorizontalHeaderLabels(["출력 위치 (A=1)", "컬럼명", "값 방식", "고정값", "참조 시트 (빈칸=현재)", "참조 셀", "인식값", "변경 후 값", "값 매핑"])
        self.added_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.added_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        layout.addWidget(self.added_table, 1)
        self.added_status = QLabel()
        self.added_status.setWordWrap(True)
        layout.addWidget(self.added_status)
        self.added_table.itemChanged.connect(self.invalidate)
        tabs.addTab(widget, "컬럼 추가")

    def _added_columns(self):
        return [AddedColumn(
            name=self.added_table.item(row, 1).text().strip(),
            position=self.added_table.cellWidget(row, 0).value(),
            mode=self.added_table.cellWidget(row, 2).currentData(),
            value=self.added_table.item(row, 3).text(),
            sheet_name=self.added_table.item(row, 4).text().strip(),
            cell=self.added_table.item(row, 5).text().strip().upper(),
            value_mapping=deepcopy(self.added_table.item(row, 1).data(Qt.UserRole) or {}),
            branch_lookup=deepcopy(self.added_table.item(row, 1).data(Qt.UserRole + 1) or {}),
        ) for row in range(self.added_table.rowCount())]

    def _set_added_columns(self, columns):
        was_loading = self.loading
        self.loading = True
        self.added_table.setRowCount(0)
        self.added_table.setRowCount(len(columns))
        for row, column in enumerate(columns):
            position = self._spin(1, 16_384, column.position)
            position.setAlignment(Qt.AlignCenter)
            self.added_table.setCellWidget(row, 0, position)
            mode = self._combo([("고정값", "fixed"), ("원본 셀 값", "cell")])
            mode.setCurrentIndex(mode.findData(column.mode))
            self.added_table.setCellWidget(row, 2, mode)
            for index, value in ((1, column.name), (3, column.value), (4, column.sheet_name), (5, column.cell)):
                item = QTableWidgetItem(value)
                item.setTextAlignment(Qt.AlignCenter)
                self.added_table.setItem(row, index, item)
            self.added_table.item(row, 1).setData(Qt.UserRole, deepcopy(column.value_mapping))
            self.added_table.item(row, 1).setData(Qt.UserRole + 1, deepcopy(column.branch_lookup))
            for index in (6, 7):
                item = QTableWidgetItem()
                item.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
                item.setTextAlignment(Qt.AlignCenter)
                self.added_table.setItem(row, index, item)
            mapping_button = QPushButton("매핑 설정")
            mapping_button.clicked.connect(lambda _checked=False, r=row: self.edit_value_mapping(r))
            self.added_table.setCellWidget(row, 8, mapping_button)
            position.valueChanged.connect(self.invalidate)
            mode.currentIndexChanged.connect(self.invalidate)
        self.loading = was_loading
        self.invalidate()

    def edit_value_mapping(self, row):
        if self.current_index < 0:
            return
        try:
            template = self.editor_template()
            template.validate()
            column = template.added_columns[row]
            book = self.entries[self.current_index]["book"]
            source = next((sheet for sheet in book.sheets if sheet.name == template.sheet_name), None)
            if source is None:
                raise ValueError("기준 시트를 선택해 주세요.")
            raw = added_column_raw_value(book, source, column)
        except ValueError as exc:
            self._error(str(exc))
            return
        key = "" if raw is None else str(raw)
        dialog = QDialog(self)
        dialog.setWindowTitle(f"{column.name} · 값 매핑")
        dialog.resize(620, 420)
        layout = QVBoxLayout(dialog)
        label = QLabel(f"현재 인식값: {key if key else '(빈값)'}\n지점 관리에서 인식값과 일치하는 필드로 지점을 찾고, 그 지점의 다른 필드 값을 가져옵니다.")
        label.setWordWrap(True)
        layout.addWidget(label)
        try:
            branch_data = load_branch_data()
        except ValueError as exc:
            self._error(str(exc))
            return
        enabled = QCheckBox("지점 관리 조회로 값 치환")
        enabled.setChecked(True)
        layout.addWidget(enabled)
        match_combo = self._combo([(label, key) for key, label in branch_data[1].items()])
        return_combo = self._combo([("가져올 필드를 선택하세요", "")] + [(label, key) for key, label in branch_data[1].items()])
        match_combo.setObjectName("branch_match_field")
        return_combo.setObjectName("branch_return_field")
        match_combo.setCurrentIndex(match_combo.findData(column.branch_lookup.get("match_field", "organization_code")))
        return_combo.setCurrentIndex(return_combo.findData(column.branch_lookup.get("return_field", "")))
        form = QFormLayout()
        form.addRow("조회할 필드 (인식값과 비교)", match_combo)
        form.addRow("가져올 필드 (출력값)", return_combo)
        layout.addLayout(form)
        result_label = QLabel()
        result_label.setWordWrap(True)
        layout.addWidget(result_label)

        def settings():
            return {"match_field": match_combo.currentData(), "return_field": return_combo.currentData()}

        def check_lookup():
            match_combo.setEnabled(enabled.isChecked())
            return_combo.setEnabled(enabled.isChecked())
            if not enabled.isChecked():
                result_label.setText("치환 해제: 인식값을 그대로 출력합니다.")
                return True
            if not match_combo.currentData() or not return_combo.currentData():
                result_label.setText("조회할 필드와 가져올 필드를 선택해 주세요.")
                return False
            try:
                branch, value = lookup_branch(raw, settings(), branch_data)
                result_label.setText(f"인식된 지점: {branch.get('branch_name', '')}\n가져올 값: {value if value is not None else '(빈값)'}")
                return True
            except ValueError as exc:
                result_label.setText(str(exc))
                return False

        match_combo.currentIndexChanged.connect(check_lookup)
        return_combo.currentIndexChanged.connect(check_lookup)
        enabled.toggled.connect(check_lookup)
        check_lookup()
        error = QLabel()
        layout.addWidget(error)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        layout.addWidget(buttons)

        def save():
            if not check_lookup():
                error.setText("조회 결과를 확인한 뒤 저장해 주세요.")
                return
            self.added_table.blockSignals(True)
            self.added_table.item(row, 1).setData(Qt.UserRole, {})
            self.added_table.item(row, 1).setData(Qt.UserRole + 1, settings() if enabled.isChecked() else {})
            self.added_table.blockSignals(False)
            self.invalidate()
            dialog.accept()

        buttons.accepted.connect(save)
        buttons.rejected.connect(dialog.reject)
        dialog.exec()

    def add_column(self, after=False):
        if self.current_index < 0 or not self.columns_table.rowCount():
            self.added_status.setText("파일을 추가하고 양식 편집에서 헤더 컬럼을 먼저 불러오세요.")
            return
        columns = self._added_columns()
        index = self.added_preview.currentIndex().column()
        position = index + 1 + int(after) if index >= 0 else 1
        for column in columns:
            if column.position >= position:
                column.position += 1
        used = {column.name for column in columns} | {column.output_name for column in self.editor_template().columns}
        number = 1
        while f"추가 컬럼 {number}" in used:
            number += 1
        columns.append(AddedColumn(f"추가 컬럼 {number}", position))
        self._set_added_columns(columns)
        self.added_table.selectRow(len(columns) - 1)
        self.added_preview.selectColumn(position - 1)

    def remove_added_column(self):
        row = self.added_table.currentRow()
        columns = self._added_columns()
        if row < 0:
            return
        removed = columns.pop(row)
        for column in columns:
            if column.position > removed.position:
                column.position -= 1
        self._set_added_columns(columns)

    def refresh_added_preview(self):
        self._set_preview(self.added_preview)
        self.added_table.blockSignals(True)
        for row in range(self.added_table.rowCount()):
            for col in (6, 7):
                self.added_table.item(row, col).setText("")
        self.added_table.blockSignals(False)
        self.added_status.clear()
        if self.current_index < 0:
            return
        template = self.editor_template()
        if not template.columns:
            self.added_status.setText("양식 편집에서 헤더 컬럼을 먼저 불러오세요.")
            return
        try:
            template.validate()
            book = self.entries[self.current_index]["book"]
            source = next((sheet for sheet in book.sheets if sheet.name == template.sheet_name), None)
            if source is None:
                raise ValueError("기준 시트를 선택해 주세요.")
            for index, column in enumerate(template.added_columns):
                raw = added_column_raw_value(book, source, column)
                mapped = added_column_value(book, source, column)
                self.added_table.blockSignals(True)
                self.added_table.item(index, 6).setText("" if raw is None else str(raw))
                self.added_table.item(index, 7).setText("" if mapped is None else str(mapped))
                self.added_table.blockSignals(False)
            selected = [column for column in template.columns if column.enabled]
            names = [column.output_name for column in selected]
            # A small layout sample; the conversion preview provides the full extraction result.
            rows = [[row[column.source_column - 1] if column.source_column <= len(row) else None
                     for column in selected]
                    for row in source.rows[template.data_start_row - 1:template.data_start_row + 19]]
            for column in sorted(template.added_columns, key=lambda column: column.position):
                value = added_column_value(book, source, column)
                names.insert(column.position - 1, column.name)
                for row in rows:
                    row.insert(column.position - 1, value)
            self._set_preview(self.added_preview, rows, [f"{get_column_letter(i + 1)} · {name}" for i, name in enumerate(names)])
            self.added_status.setText("배치 예시: 원본 시작행부터 최대 20행. 실제 추출 결과는 아래 ‘Dataset 결과’에서 자동 갱신됩니다.")
        except ValueError as exc:
            self.added_status.setText(str(exc))

    @staticmethod
    def _spin(minimum, maximum, value):
        spin = QSpinBox()
        spin.setRange(minimum, maximum)
        spin.setValue(value)
        return spin

    @staticmethod
    def _combo(options):
        combo = QComboBox()
        for label, value in options:
            combo.addItem(label, value)
        return combo

    @staticmethod
    def _view():
        view = QTableView()
        view.setEditTriggers(QAbstractItemView.NoEditTriggers)
        view.setModel(PreviewModel(parent=view))
        return view

    def _preview_tab(self, tabs, title):
        widget = QWidget()
        layout = QVBoxLayout(widget)
        combo = QComboBox()
        view = self._view()
        layout.addWidget(combo)
        layout.addWidget(view)
        tabs.addTab(widget, title)
        return combo, view

    @staticmethod
    def _set_preview(view, rows=(), columns=(), row_numbers=None):
        previous = view.model()
        view.setModel(PreviewModel(rows, columns, row_numbers, view))
        if previous:
            previous.deleteLater()

    def _error(self, text):
        self.log_view.appendPlainText(text)
        self.status.setText(text)
        QMessageBox.warning(self, "엑셀 데이터 정렬", text)

    def _run(self, action, completed, status, quiet=False):
        if self.worker is not None:
            return
        self.body.setEnabled(False)
        self.status.setText(status)
        self.worker = ExcelWorker(action, self)
        self.worker.succeeded.connect(completed)
        self.worker.failed.connect(self._preview_error if quiet else self._error)
        self.worker.finished.connect(self._finished)
        self.worker.start()

    def _finished(self):
        self.worker.deleteLater()
        self.worker = None
        self.body.setEnabled(True)

    def closeEvent(self, event):
        if self.worker is not None:
            event.ignore()
            return
        self.preview_timer.stop()
        super().closeEvent(event)

    def refresh_templates(self):
        self.templates, errors = self.store.load_all()
        self.template_list.clear()
        self.template_list.addItems([template.name for template in self.templates])
        for error in errors:
            self.log_view.appendPlainText(error)

    def add_files(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "엑셀 파일 선택", "", "Excel (*.xlsx *.xls)")
        self.load_paths(paths)

    def load_paths(self, paths):
        self._capture_current()
        existing = {entry["book"].path.resolve() for entry in self.entries}
        paths = list(dict.fromkeys(Path(path).resolve() for path in paths if Path(path).resolve() not in existing))
        if not paths:
            return
        templates = deepcopy(self.templates)

        def read():
            loaded, errors = [], []
            for path in paths:
                try:
                    book = read_workbook(path)
                    match = self.recognizer.recognize(book, templates)
                    loaded.append({"book": book, "template": deepcopy(match.selected), "status": match.reason, "recognition_log": recognition_log(match)})
                except Exception as exc:
                    errors.append(f"{path.name}: {exc}")
            return loaded, errors

        def complete(result):
            loaded, errors = result
            self.entries.extend(loaded)
            self._refresh_files()
            for entry in loaded:
                self.log_view.appendPlainText(f"{entry['book'].path.name}: {entry['recognition_log']}")
            self.status.setText(f"{len(loaded)}개 파일을 불러왔습니다. 파일별 양식을 확인해 주세요.")
            if errors:
                self._error("불러오기 실패:\n" + "\n".join(errors))

        self._run(read, complete, "엑셀 파일 읽기 및 양식 자동 인식 중…")

    def _capture_current(self):
        if 0 <= self.current_index < len(self.entries):
            draft = self.editor_template()
            if draft.columns:
                self.entries[self.current_index]["template"] = draft

    def _refresh_files(self):
        self.files.blockSignals(True)
        self.files.clear()
        for entry in self.entries:
            template = entry["template"]
            self.files.addItem(f"{entry['book'].path.name}  |  {template.name if template else '양식 선택 필요'}")
            self.files.item(self.files.count() - 1).setToolTip(str(entry["book"].path) + "\n" + entry["status"])
        index = min(max(self.current_index, 0), len(self.entries) - 1)
        self.files.setCurrentRow(index)
        self.files.blockSignals(False)
        self.current_index = -1
        self._file_changed(index)

    def _file_changed(self, index):
        self._capture_current()
        self.current_index = index
        if index < 0:
            self.invalidate()
            self._set_preview(self.raw_table)
            return
        entry = self.entries[index]
        book = entry["book"]
        self.load_editor(entry["template"] or Template("새 양식", book.sheets[0].name))
        self.status.setText(entry["status"])

    def remove_file(self):
        if self.current_index >= 0:
            del self.entries[self.current_index]
            self.current_index = -1
            self._refresh_files()

    def _header_changed(self, value):
        if self.start_spin.value() <= value:
            self.start_spin.setValue(value + 1)

    def load_editor(self, template):
        self.loading = True
        self.name_edit.setText(template.name)
        self.sheet_combo.clear()
        self.sheet_combo.addItems([sheet.name for sheet in self.entries[self.current_index]["book"].sheets])
        index = self.sheet_combo.findText(template.sheet_name)
        if index < 0:
            self.sheet_combo.addItem(template.sheet_name)
            index = self.sheet_combo.count() - 1
        self.sheet_combo.setCurrentIndex(index)
        self.header_spin.setValue(template.header_row)
        self.start_spin.setValue(template.data_start_row)
        self.end_spin.setValue(template.data_end_row)
        self.first_column_spin.setValue(template.header_start_column)
        self.last_column_spin.setValue(template.header_end_column)
        for widget, value in (
            (self.merge_combo, template.merge_mode), (self.table_combo, template.table_mode),
            (self.sheet_mode_combo, template.sheet_mode), (self.match_combo, template.header_match),
        ):
            widget.setCurrentIndex(widget.findData(value))
        for widget, value in (
            (self.hidden_rows, template.include_hidden_rows), (self.hidden_columns, template.include_hidden_columns),
            (self.hidden_sheets, template.include_hidden_sheets), (self.blank_stop, template.stop_at_blank),
            (self.source_check, template.include_source),
        ):
            widget.setChecked(value)
        self.stop_edit.setText(", ".join(template.stop_values))
        self._set_columns(template.columns)
        self._set_added_columns(template.added_columns)
        self.loading = False
        self.invalidate()
        self.show_raw()

    def _set_columns(self, columns):
        self.columns_table.blockSignals(True)
        self.columns_table.setRowCount(len(columns))
        for row, column in enumerate(columns):
            check = QTableWidgetItem()
            check.setFlags(Qt.ItemIsEnabled | Qt.ItemIsUserCheckable)
            check.setCheckState(Qt.Checked if column.enabled else Qt.Unchecked)
            self.columns_table.setItem(row, 0, check)
            position = QTableWidgetItem(get_column_letter(column.source_column))
            position.setData(Qt.UserRole, column.source_column)
            position.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            self.columns_table.setItem(row, 1, position)
            header = QTableWidgetItem(column.header)
            header.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable)
            self.columns_table.setItem(row, 2, header)
            self.columns_table.setItem(row, 3, QTableWidgetItem(column.output_name))
            for index in range(self.columns_table.columnCount()):
                self.columns_table.item(row, index).setTextAlignment(Qt.AlignCenter)
        self.columns_table.blockSignals(False)
        self.invalidate()

    def editor_template(self):
        columns = [ColumnRule(
            self.columns_table.item(row, 1).data(Qt.UserRole),
            self.columns_table.item(row, 2).text(), self.columns_table.item(row, 3).text().strip(),
            self.columns_table.item(row, 0).checkState() == Qt.Checked,
        ) for row in range(self.columns_table.rowCount())]
        prior = self.entries[self.current_index]["template"] if self.current_index >= 0 else None
        return Template(
            name=self.name_edit.text().strip(), sheet_name=self.sheet_combo.currentText(),
            header_row=self.header_spin.value(), data_start_row=self.start_spin.value(),
            data_end_row=self.end_spin.value(), columns=columns,
            added_columns=self._added_columns(),
            header_start_column=self.first_column_spin.value(), header_end_column=self.last_column_spin.value(),
            merge_mode=self.merge_combo.currentData(), table_mode=self.table_combo.currentData(),
            sheet_mode=self.sheet_mode_combo.currentData(), header_match=self.match_combo.currentData(),
            include_hidden_rows=self.hidden_rows.isChecked(), include_hidden_columns=self.hidden_columns.isChecked(),
            include_hidden_sheets=self.hidden_sheets.isChecked(), stop_at_blank=self.blank_stop.isChecked(),
            stop_values=[value.strip() for value in self.stop_edit.text().split(",") if value.strip()],
            include_source=self.source_check.isChecked(), signature=deepcopy(prior.signature) if prior else {},
        )

    def read_columns(self):
        if self.current_index < 0:
            return
        template = self.editor_template()
        book = self.entries[self.current_index]["book"]
        source = next((sheet for sheet in book.sheets if sheet.name == template.sheet_name), None)
        if source is None:
            self._error("이 파일에 기준 시트가 없습니다. 실제 시트를 선택해 주세요.")
            return
        sheet = normalize(source, template)
        if template.header_row not in sheet.row_numbers:
            self._error("지정한 헤더행이 비어 있거나 숨김 처리되었습니다.")
            return
        row = sheet.rows[sheet.row_numbers.index(template.header_row)]
        columns, names = [], set()
        for position, value in zip(sheet.column_numbers, row):
            if position < template.header_start_column or template.header_end_column and position > template.header_end_column:
                continue
            if is_empty(value):
                continue
            header = str(value).strip()
            output = header if header not in names else f"{header}_{get_column_letter(position)}"
            names.add(output)
            columns.append(ColumnRule(position, header, output))
        self._set_columns(columns)
        self.status.setText("컬럼을 불러왔습니다. 사용할 컬럼과 출력 이름을 지정한 뒤 미리보기를 확인해 주세요.")

    def _selected_template(self):
        index = self.template_list.currentRow()
        if index < 0 or self.current_index < 0:
            self._error("파일과 등록된 양식을 선택해 주세요.")
            return None
        return deepcopy(self.templates[index])

    def apply_template(self):
        template = self._selected_template()
        if template is not None:
            self.entries[self.current_index]["template"] = template
            self.entries[self.current_index]["status"] = f"수동 선택: {template.name}"
            self._refresh_files()

    def apply_template_all(self):
        template = self._selected_template()
        if template is not None:
            for entry in self.entries:
                entry["template"] = deepcopy(template)
                entry["status"] = f"수동 일괄 선택: {template.name} (미리보기로 호환 여부 확인)"
            self._refresh_files()

    def new_template(self):
        if self.current_index >= 0:
            entry = self.entries[self.current_index]
            entry["template"] = None
            self.load_editor(Template("새 양식", entry["book"].sheets[0].name))
            self.status.setText("새 양식: 이름과 기준 헤더를 지정해 주세요.")

    def recognize_all(self):
        if not self.entries:
            return
        books = [entry["book"] for entry in self.entries]
        templates = deepcopy(self.templates)

        def complete(matches):
            for entry, match in zip(self.entries, matches):
                entry["template"] = deepcopy(match.selected)
                entry["status"] = match.reason
                entry["recognition_log"] = recognition_log(match)
                self.log_view.appendPlainText(f"{entry['book'].path.name}: {entry['recognition_log']}")
            self._refresh_files()

        self._run(lambda: [self.recognizer.recognize(book, templates) for book in books], complete, "양식 자동 인식 중…")

    def save_template(self):
        if self.current_index < 0:
            return
        try:
            template = self.editor_template()
            template.validate()
            template.signature = build_signature(self.entries[self.current_index]["book"], template)
            path = self.store.directory / f"{template.name}.json"
            if path.exists() and QMessageBox.question(
                self, "양식 덮어쓰기", f"'{template.name}' 양식을 현재 설정으로 바꾸시겠습니까?",
                QMessageBox.Yes | QMessageBox.No,
            ) != QMessageBox.Yes:
                return
            path = self.store.save(template)
            self.entries[self.current_index]["template"] = template
            self.entries[self.current_index]["status"] = f"저장된 양식: {template.name}"
            self.refresh_templates()
            self._refresh_files()
            self.status.setText(f"양식 저장 완료: {path}")
        except (ValueError, OSError) as exc:
            self._error(str(exc))

    def invalidate(self, *_args):
        if self.loading:
            return
        self.result = None
        self.dataset_combo.clear()
        self._set_preview(self.dataset_table)
        self.refresh_added_preview()
        self.dataset_hint.setText("설정 변경 반영 중…")
        self.preview_timer.start()

    def _preview_error(self, message):
        self.dataset_hint.setText(f"Dataset 미리보기: {message}")
        if message != self.last_preview_error:
            self.log_view.appendPlainText(f"Dataset 미리보기: {message}")
            self.last_preview_error = message

    def _auto_preview(self):
        if self.worker is not None:
            self.preview_timer.start()
            return
        if self.current_index < 0 or not self.columns_table.rowCount():
            self.dataset_hint.setText("파일을 선택하고 헤더 컬럼을 불러오세요.")
            return
        self.preview(automatic=True)

    def show_raw(self, *_args):
        if self.loading or self.current_index < 0:
            return
        book = self.entries[self.current_index]["book"]
        source = next((sheet for sheet in book.sheets if sheet.name == self.sheet_combo.currentText()), None)
        if source:
            width = max((len(row) for row in source.rows), default=0)
            rows = [row + [None] * (width - len(row)) for row in source.rows[:500]]
            self._set_preview(self.raw_table, rows, [get_column_letter(index + 1) for index in range(width)])
        else:
            self._set_preview(self.raw_table)

    def preview(self, automatic=False):
        self.preview_timer.stop()
        if self.current_index < 0:
            return
        try:
            template = self.editor_template()
            template.validate()
        except ValueError as exc:
            (self._preview_error if automatic else self._error)(str(exc))
            return
        self.entries[self.current_index]["template"] = template
        book = self.entries[self.current_index]["book"]

        def complete(result):
            self._display_result(result)
            for warning in result.warnings:
                self.log_view.appendPlainText(f"{book.path.name}: {warning}")
            self.status.setText(f"변환 완료: {sum(len(dataset.rows) for dataset in result.datasets):,}행 / {len(result.datasets)}개 Dataset")

        self._run(lambda: convert(book, template), complete, "Dataset 갱신 중…", quiet=automatic)

    def _display_result(self, result):
        self.dataset_combo.clear()
        self.result = result
        self.dataset_combo.addItems([f"{dataset.name} ({len(dataset.rows):,}행)" for dataset in result.datasets])
        self.show_dataset()
        self.last_preview_error = ""
        self.dataset_hint.setText(f"현재 설정 적용 완료: {sum(len(dataset.rows) for dataset in result.datasets):,}행 (미리보기 최대 500행)")

    def show_dataset(self, *_args):
        index = self.dataset_combo.currentIndex()
        if self.result and index >= 0:
            dataset = self.result.datasets[index]
            self._set_preview(self.dataset_table, dataset.rows, dataset.columns)

    def export_all(self):
        if not self.entries:
            return
        self._capture_current()
        for entry in self.entries:
            if entry["template"] is None:
                self._error(f"{entry['book'].path.name}: 양식을 먼저 선택해 주세요.")
                return
            try:
                entry["template"].validate()
            except ValueError as exc:
                self._error(f"{entry['book'].path.name}: {exc}")
                return
        path, _ = QFileDialog.getSaveFileName(self, "Dataset Excel 저장", "Dataset.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"
        jobs = [(entry["book"], deepcopy(entry["template"])) for entry in self.entries]

        def write():
            datasets, warnings = convert_many(jobs)
            export_excel(datasets, path, [book.path for book, _template in jobs])
            return sum(len(dataset.rows) for dataset in datasets), warnings

        def complete(result):
            count, warnings = result
            for warning in warnings:
                self.log_view.appendPlainText(warning)
            self.status.setText(f"Excel 저장 완료: {path} / {count:,}행")
            self.log_view.appendPlainText(self.status.text())

        self._run(write, complete, "전체 파일을 변환하고 Excel에 저장하는 중…")
