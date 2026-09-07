import json
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import Workbook, load_workbook


ATTENDANCE_HEADERS = (
    "지점",
    "직원명",
    "직종",
    "생년월일",
    "날짜",
    "출근시간",
    "퇴근시간",
    "휴게시간",
    "근무유형",
)
LABEL_EXTENSIONS = {".json", ".xlsx"}


def _cell_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return value


def read_excel_label(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError(f"비어 있는 엑셀 정답 파일입니다: {path}") from exc

        headers = [str(value).strip() if value is not None else "" for value in raw_headers]
        if not any(headers):
            raise ValueError(f"엑셀 정답 파일의 첫 행에 열 이름이 없습니다: {path}")
        if len(set(headers)) != len(headers) or "" in headers:
            raise ValueError(f"엑셀 정답 파일의 열 이름은 비어 있거나 중복될 수 없습니다: {path}")

        records = []
        for row in rows:
            values = [_cell_value(value) for value in row[:len(headers)]]
            values.extend([None] * (len(headers) - len(values)))
            if not any(value is not None and value != "" for value in values):
                continue
            records.append(dict(zip(headers, values)))
        return {"columns": headers, "rows": records}
    finally:
        workbook.close()


def read_label(path):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        return read_excel_label(path)
    if path.suffix.lower() == ".json":
        with path.open("r", encoding="utf-8") as file:
            return json.load(file)
    raise ValueError(f"지원하지 않는 정답 파일 형식입니다: {path.suffix}")


def _table_from_result(result):
    if isinstance(result, dict) and isinstance(result.get("rows"), list):
        rows = result["rows"]
        columns = result.get("columns") or []
    elif isinstance(result, list):
        rows = result
        columns = []
    else:
        raise ValueError("분석 결과가 엑셀 표 형태가 아닙니다.")
    if not all(isinstance(row, dict) for row in rows):
        raise ValueError("분석 결과의 각 행은 열 이름과 값으로 구성되어야 합니다.")
    if not columns:
        columns = list(ATTENDANCE_HEADERS)
        for row in rows:
            for key in row:
                if key not in columns:
                    columns.append(key)
    return columns, rows


def write_excel_result(result, path):
    columns, rows = _table_from_result(result)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "출근부 분석"
    worksheet.append(columns)
    for row in rows:
        worksheet.append([row.get(column) for column in columns])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(columns, 1):
        values = [str(column)] + [str(row.get(column, "") or "") for row in rows]
        worksheet.column_dimensions[worksheet.cell(1, index).column_letter].width = min(
            max(len(value) for value in values) + 2, 30
        )
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
